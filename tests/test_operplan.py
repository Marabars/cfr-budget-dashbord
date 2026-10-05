"""Тесты вкладки «Оперплан»: импорт листа, слои, повтор года, замок, правка.

Файлы-фикстуры собираются на месте — тесты не зависят от рабочего файла
ликвидности на сетевом диске.
"""
from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import db, operplan, repo
from app.importers import cfo_budgets
from app.main import app

MONTH_NAMES = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
               "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]

IT = "Первая дирекция"
HR = "Вторая дирекция"


***REMOVED*** ---------- фикстуры ----------

def make_budgets_sheet() -> bytes:
    """Лист «Бюджеты ЦФО» в раскладке рабочего файла: шапка в 3-й строке."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Бюджеты ЦФО"

    header = ["Статья оборотов", "ЦФО", "Контрагент", "Договор", "Детализация",
              "Комментарий", "ИТОГО"] + MONTH_NAMES
    for col, value in enumerate(header, start=1):
        ws.cell(row=3, column=col, value=value)

    def put(row, values, months, total=None):
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        ws.cell(row=row, column=7, value=sum(months) if total is None else total)
        for m, amount in enumerate(months):
            ws.cell(row=row, column=8 + m, value=amount or None)

    ***REMOVED*** блок ИТ: два бюджета, детализация заполнена
    put(4, ["ИТОГО", IT, None, "Бюджет обязательных затрат"],
        [100_000, 50_000] + [0] * 10, total=150_000)
    put(5, ["4411_Оборудование", IT, "Астра ООО", "Договор №1", "ИТ инфраструктура",
            "Серверы"], [100_000] + [0] * 11)
    put(6, ["4411_Оборудование", IT, None, None, None, "Картриджи"],
        [0, 50_000] + [0] * 10)

    put(7, ["ИТОГО", IT, None, "Бюджет развития"], [0, 0, 200_000] + [0] * 9,
        total=200_000)
    put(8, ["4430_Разработка", IT, "Софт ООО", "Договор №2", "Платформа", "Доработки"],
        [0, 0, 200_000] + [0] * 9)

    ***REMOVED*** блок ДУП: детализации нет, только комментарий; одна строка с чужой подписью ЦФО
    put(9, ["ИТОГО", HR, None, None], [15_000] + [10_000] * 10 + [60_000], total=175_000)
    put(10, ["4109_Обучение", HR, None, None, None, "Английский"], [0] * 11 + [30_000])
    put(11, ["4109_Обучение", HR, None, None, None, "Библиотека"], [0] * 11 + [20_000])
    put(12, ["4110_Соцрасходы", HR, "Подрядчик ООО", "Договор №3", None, "Подарки"],
        [10_000] * 12)
    put(13, ["4110_Соцрасходы", "Эпсилон", None, None, None, "Через другое ЮЛ"],
        [5_000] + [0] * 11)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def make_articles_sheet() -> bytes:
    """Лист «ЦФО Статьи»: шапка смещена вправо и вниз, название ЦФО тянется по блоку."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ЦФО Статьи"
    ws.cell(row=2, column=2, value="Соответсвие статей ЦФО")
    ws.cell(row=4, column=2, value="ЦФО -Держатель бюджета ")
    ws.cell(row=4, column=3, value="Статьи управленческого учета, характерные для ЦФО")
    ws.cell(row=5, column=2, value=HR)
    ws.cell(row=5, column=3, value="4109_Обучение")
    ws.cell(row=6, column=3, value="4110_Соцрасходы")        ***REMOVED*** ЦФО тянется сверху
    ws.cell(row=7, column=3, value="4114_Командировочные")   ***REMOVED*** статьи нет в строках бюджета
    ws.cell(row=8, column=2, value="Бухгалтерия")            ***REMOVED*** такого ЦФО в бюджетах нет
    ws.cell(row=8, column=3, value="4106_Аудит")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("BUDGET_DB", str(tmp_path / "test.db"))
    db.set_db_path(tmp_path / "test.db")
    db.init_db()
    db.set_setting("year", "2026")
    db.set_setting("operplan_year", "2027")
    db.set_setting("operplan_base_year", "2026")
    with TestClient(app) as c:
        yield c


@pytest.fixture()
def loaded(client):
    """Слой планирования с загруженным базовым 2026 годом."""
    result = operplan.import_budgets(make_budgets_sheet(), "бюджеты.xlsx", 2026, "Тест")
    assert result["rows"] == 7
    return result


@pytest.fixture()
def planned(loaded):
    """План 2027, наполненный повтором базы: правят только планируемый год."""
    return operplan.copy_year("Тест", 2026, 2027)


def plan_row(cfo_name: str | None = None) -> dict:
    """Строка плана 2027 — при необходимости конкретной дирекции."""
    rows = repo.list_rows(2027, layer=operplan.LAYER)
    if cfo_name:
        rows = [r for r in rows if r["cfo_name"] == cfo_name]
    assert rows, "в плане нет подходящей строки"
    return rows[0]


***REMOVED*** ---------- разбор листа ----------

def test_держатель_бюджета_берётся_из_строки_итого(loaded):
    """Строка с подписью «Эпсилон» лежит внутри блока ДУП и остаётся за ДУП."""
    rows = repo.list_rows(2026, layer=operplan.LAYER)
    gasfor = [r for r in rows if r["comment"] == "Через другое ЮЛ"]
    assert len(gasfor) == 1
    assert gasfor[0]["cfo_name"] == HR
    ***REMOVED*** прежняя подпись не теряется
    assert "Эпсилон" in gasfor[0]["approver_comment"]


def test_бюджеты_ит_разложены_по_двум_блокам(loaded):
    budgets = {b["name"] for b in repo.list_budgets(2026, layer=operplan.LAYER)}
    assert "Бюджет обязательных затрат" in budgets
    assert "Бюджет развития" in budgets


def test_контрольная_сумма_блока_сверяется():
    parsed = cfo_budgets.parse(make_budgets_sheet())
    assert cfo_budgets.check_totals(parsed) == []


def test_расхождение_итого_попадает_в_отчёт():
    """ИТОГО блока занижено — импорт обязан об этом сказать, а не молча сойтись."""
    wb = openpyxl.load_workbook(io.BytesIO(make_budgets_sheet()))
    wb["Бюджеты ЦФО"].cell(row=9, column=7, value=1)   ***REMOVED*** ИТОГО блока ДУП
    buffer = io.BytesIO()
    wb.save(buffer)
    parsed = cfo_budgets.parse(buffer.getvalue())
    mismatches = cfo_budgets.check_totals(parsed)
    assert len(mismatches) == 1
    assert mismatches[0]["cfo"] == HR
    assert mismatches[0]["declared"] == 1


def test_детализация_у_ит_из_бюджета_у_остальных_комментарий(client, loaded):
    view = operplan.view(2026)
    blocks = {b["cfo_name"] + "|" + b["budget_name"]: b for b in view["blocks"]}

    it = blocks[IT + "|Бюджет обязательных затрат"]
    assert it["detail_source"] == "detail"
    server = next(r for r in it["rows"] if r["counterparty"] == "Астра ООО")
    assert server["detail_text"] == "ИТ инфраструктура"

    hr = blocks[HR + "|Основной бюджет"]
    assert hr["detail_source"] == "comment"
    english = next(r for r in hr["rows"] if r["comment"] == "Английский")
    assert english["detail_text"] == "Английский"


def test_у_ит_пустая_детализация_подменяется_комментарием(loaded):
    """Иначе у ИТ колонка была бы пустой там, где в файле заполнен только комментарий."""
    view = operplan.view(2026)
    it = next(b for b in view["blocks"] if b["budget_name"] == "Бюджет обязательных затрат")
    cartridges = next(r for r in it["rows"] if r["comment"] == "Картриджи")
    assert cartridges["detail"] == ""
    assert cartridges["detail_text"] == "Картриджи"


***REMOVED*** ---------- слои не смешиваются ----------

def test_слой_планирования_не_виден_дашборду(client, loaded):
    dashboard = client.get("/api/dashboard").json()
    assert dashboard["cards"] == []
    assert dashboard["totals"]["plan_year"] == 0
    assert operplan.view(2026)["totals"]["rows_count"] == 7


def test_строки_исполнения_не_видны_оперплану(client):
    cfo_id = repo.ensure_cfo(IT)
    client.post("/api/rows", json={"cfo_id": cfo_id, "article": "4420_ПО",
                                   "plan_year": 500_000, "spread": "even"})
    assert client.get("/api/dashboard").json()["totals"]["plan_year"] == 500_000
    assert client.get("/api/operplan?year=2026").json()["totals"]["rows_count"] == 0


def test_одноимённый_бюджет_живёт_в_обоих_слоях(client, loaded):
    """UNIQUE бюджета включает слой — иначе импорт упал бы на уже занятом имени."""
    cfo_id = repo.ensure_cfo(IT)
    main_id = repo.ensure_budget(cfo_id, 2026, "Бюджет обязательных затрат")
    plan_id = repo.ensure_budget(cfo_id, 2026, "Бюджет обязательных затрат",
                                 layer=operplan.LAYER)
    assert main_id != plan_id


***REMOVED*** ---------- повтор бюджета прошлого года ----------

def test_повтор_года_переносит_все_строки_с_разбивкой(client, loaded):
    base = operplan.view(2026)["totals"]
    result = operplan.copy_year("Тест", 2026, 2027)
    assert result["copied"] == 7 and result["skipped"] == 0

    plan = operplan.view(2027)["totals"]
    assert plan["rows_count"] == 7
    assert plan["plan_year"] == pytest.approx(base["plan_year"])
    assert plan["months"] == pytest.approx(base["months"])


def test_индексация_умножает_суммы(client, loaded):
    base = operplan.view(2026)["totals"]["plan_year"]
    operplan.copy_year("Тест", 2026, 2027, index_pct=10)
    assert operplan.view(2027)["totals"]["plan_year"] == pytest.approx(base * 1.1, rel=1e-6)


def test_повторный_запуск_ничего_не_задваивает(client, loaded):
    operplan.copy_year("Тест", 2026, 2027)
    again = operplan.copy_year("Тест", 2026, 2027)
    assert again["copied"] == 0
    assert again["skipped"] == 7
    assert operplan.view(2027)["totals"]["rows_count"] == 7


def test_строки_с_одинаковым_ключом_не_схлопываются(client, loaded):
    """У ДУП две строки одной статьи без контрагента и договора — обе обязаны переехать."""
    operplan.copy_year("Тест", 2026, 2027)
    rows = repo.list_rows(2027, layer=operplan.LAYER)
    training = [r for r in rows if r["article"] == "4109_Обучение"]
    assert len(training) == 2
    assert {r["comment"] for r in training} == {"Английский", "Библиотека"}


def test_индексация_за_пределами_разумного_отклоняется(client, loaded):
    response = client.post("/api/operplan/copy-year", json={"index_pct": -150})
    assert response.status_code == 422


def test_копирование_в_тот_же_год_отклоняется(client, loaded):
    with pytest.raises(ValueError):
        operplan.copy_year("Тест", 2026, 2026)


***REMOVED*** ---------- были в базовом году, нет в плане ----------

def test_пропавшие_строки_находятся_по_статье_контрагенту_договору(client, loaded):
    operplan.copy_year("Тест", 2026, 2027)
    assert operplan.missing_rows(2027)["count"] == 0

    rows = repo.list_rows(2027, layer=operplan.LAYER)
    victim = next(r for r in rows if r["article"] == "4430_Разработка")
    repo.delete_row(int(victim["id"]))

    missing = operplan.missing_rows(2027)
    assert missing["count"] == 1
    assert missing["items"][0]["article"] == "4430_Разработка"
    assert missing["total"] == pytest.approx(200_000)


def test_добавить_в_план_возвращает_строку(client, loaded):
    operplan.copy_year("Тест", 2026, 2027)
    rows = repo.list_rows(2027, layer=operplan.LAYER)
    repo.delete_row(int(next(r for r in rows if r["article"] == "4430_Разработка")["id"]))

    missing = operplan.missing_rows(2027)
    response = client.post(f"/api/operplan/missing/{missing['items'][0]['id']}/add?year=2027")
    assert response.status_code == 201
    assert operplan.missing_rows(2027)["count"] == 0

    restored = next(r for r in repo.list_rows(2027, layer=operplan.LAYER)
                    if r["article"] == "4430_Разработка")
    months = repo.plan_months(2027, layer=operplan.LAYER)[int(restored["id"])]
    assert months[2] == pytest.approx(200_000)


***REMOVED*** ---------- KPI ----------

def test_kpi_считает_оба_года_и_помнит_что_факта_нет(client, loaded):
    operplan.copy_year("Тест", 2026, 2027, index_pct=25)
    kpi = client.get("/api/operplan").json()["kpi"]
    assert kpi["plan_year"] == 2027 and kpi["base_year"] == 2026
    assert kpi["plan_total"] == pytest.approx(kpi["base_total"] * 1.25, rel=1e-6)
    assert kpi["base_ratio"] == pytest.approx(1 / 1.25, rel=1e-6)
    assert kpi["fact_loaded"] is False
    assert kpi["base_fact"] == 0


***REMOVED*** ---------- замок ----------

def test_замок_закрывает_правку_руководителю_но_не_финотделу(client, planned):
    cfo_id = repo.ensure_cfo(HR)
    row_id = int(plan_row(HR)["id"])

    assert client.post("/api/operplan/lock",
                       json={"year": 2027, "locked": True}).status_code == 200

    blocked = client.patch(f"/api/rows/{row_id}", json={"counterparty": "Нельзя"},
                           headers={"X-User": f"cfo:{cfo_id}"})
    assert blocked.status_code == 403
    assert "заблокировано" in blocked.json()["detail"]

    assert client.patch(f"/api/rows/{row_id}", json={"counterparty": "Можно"}).status_code == 200


def test_замок_ставит_только_финотдел(client, loaded):
    cfo_id = repo.ensure_cfo(HR)
    response = client.post("/api/operplan/lock", json={"year": 2026, "locked": True},
                           headers={"X-User": f"cfo:{cfo_id}"})
    assert response.status_code == 403


def test_замок_виден_в_ответе_вкладки(client, loaded):
    cfo_id = repo.ensure_cfo(HR)
    client.post("/api/operplan/lock", json={"year": 2026, "locked": True})
    view = client.get("/api/operplan?year=2026", headers={"X-User": f"cfo:{cfo_id}"}).json()
    assert view["locked"] is True
    assert view["can_edit"] is False
    assert client.get("/api/operplan?year=2026").json()["can_edit"] is True


def test_удаление_строки_при_замке_закрыто_руководителю(client, loaded):
    cfo_id = repo.ensure_cfo(HR)
    row_id = int(next(r for r in repo.list_rows(2026, layer=operplan.LAYER)
                      if r["cfo_id"] == cfo_id)["id"])
    client.post("/api/operplan/lock", json={"year": 2026, "locked": True})
    assert client.delete(f"/api/rows/{row_id}",
                         headers={"X-User": f"cfo:{cfo_id}"}).status_code == 403


def test_руководитель_правит_и_удаляет_свою_строку_плана(client, planned):
    """Без замка руководитель ЦФО редактирует план своего ЦФО."""
    cfo_id = repo.ensure_cfo(HR)
    head = {"X-User": f"cfo:{cfo_id}"}
    row_id = int(plan_row(HR)["id"])

    assert client.get("/api/operplan?year=2027", headers=head).json()["can_edit"] is True

    patched = client.patch(f"/api/rows/{row_id}", json={"counterparty": "Свой контрагент"},
                           headers=head)
    assert patched.status_code == 200
    assert patched.json()["row"]["counterparty"] == "Свой контрагент"

    assert client.delete(f"/api/rows/{row_id}", headers=head).status_code == 200
    assert repo.get_row(row_id) is None


def test_kpi_показывает_насколько_набран_план_от_базы(client, loaded):
    """Шкала первой плитки растёт по мере наполнения плана и не делится на ноль."""
    assert client.get("/api/operplan").json()["kpi"]["plan_ratio"] == 0.0
    operplan.copy_year("Тест", 2026, 2027, index_pct=25)
    assert client.get("/api/operplan").json()["kpi"]["plan_ratio"] == pytest.approx(1.25,
                                                                                    rel=1e-6)


def test_руководитель_не_правит_чужое_цфо(client, loaded):
    hr_id = repo.ensure_cfo(HR)
    it_row = int(next(r for r in repo.list_rows(2026, layer=operplan.LAYER)
                      if r["cfo_name"] == IT)["id"])
    response = client.patch(f"/api/rows/{it_row}", json={"counterparty": "Чужое"},
                            headers={"X-User": f"cfo:{hr_id}"})
    assert response.status_code == 403


***REMOVED*** ---------- прошлый год не правится ----------

def test_прошлый_год_не_правится_по_ячейке(client, loaded):
    """База для сравнения: ни правки, ни удаления — даже финотделу."""
    row_id = int(repo.list_rows(2026, layer=operplan.LAYER)[0]["id"])

    patched = client.patch(f"/api/rows/{row_id}", json={"counterparty": "Нельзя"})
    assert patched.status_code == 403
    assert "база для сравнения" in patched.json()["detail"]

    deleted = client.delete(f"/api/rows/{row_id}")
    assert deleted.status_code == 403
    assert repo.get_row(row_id) is not None, "строка должна остаться"


def test_строку_в_прошлый_год_не_завести(client, loaded):
    response = client.post("/api/operplan/rows?year=2026",
                           json={"cfo_id": repo.ensure_cfo(HR), "article": "4109_Обучение"})
    assert response.status_code == 403


def test_прошлый_год_по_прежнему_заливается_целиком(client, loaded):
    """Запрет на правку по ячейке не мешает перезалить базу или залить факт."""
    assert client.post(
        "/api/operplan/import",
        data={"year": "2026", "kind": "plan"},
        files={"file": ("бюджеты.xlsx", make_budgets_sheet(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).status_code == 200
    assert client.post(
        "/api/operplan/import",
        data={"year": "2026", "kind": "fact"},
        files={"file": ("факт.xlsx", make_fact_sheet(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).status_code == 200


***REMOVED*** ---------- правка строки ----------

def test_правка_месяцев_пересчитывает_годовую_сумму(client, planned):
    row_id = int(plan_row()["id"])
    months = [1000.0] * 12
    response = client.patch(f"/api/rows/{row_id}", json={"months": months})
    assert response.status_code == 200
    assert response.json()["row"]["plan_year"] == pytest.approx(12_000)


def test_правка_одного_поля_не_трогает_остальные(client, planned):
    row = plan_row()
    before = dict(row)
    client.patch(f"/api/rows/{row['id']}", json={"counterparty": "Новый контрагент"})
    after = repo.get_row(int(row["id"]))
    assert after["counterparty"] == "Новый контрагент"
    assert after["article"] == before["article"]
    assert after["plan_year"] == before["plan_year"]


def test_пустая_статья_при_правке_отклоняется(client, planned):
    row_id = int(plan_row()["id"])
    assert client.patch(f"/api/rows/{row_id}", json={"article": "  "}).status_code == 422


def test_отрицательный_месяц_отклоняется(client, planned):
    row_id = int(plan_row()["id"])
    months = [-1.0] + [0.0] * 11
    assert client.patch(f"/api/rows/{row_id}", json={"months": months}).status_code == 422


def test_неполный_список_месяцев_отклоняется(client, planned):
    row_id = int(plan_row()["id"])
    assert client.patch(f"/api/rows/{row_id}", json={"months": [1, 2, 3]}).status_code == 422


def test_правка_пишется_в_журнал(client, planned):
    row_id = int(plan_row()["id"])
    client.patch(f"/api/rows/{row_id}", json={"counterparty": "Журнальный"})
    audit = client.get("/api/audit").json()
    entry = next(a for a in audit if a["action"] == "Изменение строки бюджета")
    assert "Журнальный" in entry["change"]


def test_новая_строка_плана_попадает_в_нужный_слой_и_год(client, loaded):
    cfo_id = repo.ensure_cfo(HR)
    response = client.post("/api/operplan/rows?year=2027", json={
        "cfo_id": cfo_id, "budget_name": "Основной бюджет", "article": "4114_Командировочные",
        "counterparty": "Аэрофлот", "months": [1000.0] * 12,
    })
    assert response.status_code == 201
    view = client.get("/api/operplan?year=2027").json()
    assert view["totals"]["rows_count"] == 1
    assert view["totals"]["plan_year"] == pytest.approx(12_000)
    assert client.get("/api/dashboard").json()["totals"]["plan_year"] == 0


***REMOVED*** ---------- справочники ----------

def test_справочник_статей_объединяет_лист_и_строки(client, loaded):
    operplan.import_articles(make_articles_sheet(), "статьи.xlsx", "Тест")
    refs = client.get("/api/operplan/references?year=2026").json()
    hr_id = str(repo.ensure_cfo(HR))
    articles = refs["by_cfo"][hr_id]["articles"]
    assert "4109_Обучение" in articles          ***REMOVED*** есть и в листе, и в строках
    assert "4114_Командировочные" in articles   ***REMOVED*** только в листе-справочнике
    assert "4110_Соцрасходы" in articles


def test_неизвестное_цфо_справочника_не_теряется_молча(client, loaded):
    result = operplan.import_articles(make_articles_sheet(), "статьи.xlsx", "Тест")
    assert "Бухгалтерия" in result["unknown_cfo"]


def test_договоры_сгруппированы_по_контрагенту(client, loaded):
    refs = client.get("/api/operplan/references?year=2026").json()
    it_id = str(repo.ensure_cfo(IT))
    contracts = refs["by_cfo"][it_id]["contracts"]
    assert contracts["Астра ООО"] == ["Договор №1"]
    assert contracts["Софт ООО"] == ["Договор №2"]


def test_общий_список_собирает_все_цфо(client, loaded):
    refs = client.get("/api/operplan/references?year=2026").json()
    assert "4109_Обучение" in refs["all"]["articles"]
    assert "4411_Оборудование" in refs["all"]["articles"]
    assert "Астра ООО" in refs["all"]["counterparties"]


def test_руководителю_справочник_отдаёт_только_его_цфо(client, loaded):
    """«Показать все» — привилегия финотдела: чужие контрагенты не уезжают на клиент."""
    it_id = repo.ensure_cfo(IT)
    hr_id = repo.ensure_cfo(HR)
    refs = client.get("/api/operplan/references?year=2026",
                      headers={"X-User": f"cfo:{it_id}"}).json()

    assert set(refs["by_cfo"]) == {str(it_id)}
    assert str(hr_id) not in refs["by_cfo"]
    assert "4411_Оборудование" in refs["all"]["articles"]      ***REMOVED*** своё
    assert "4109_Обучение" not in refs["all"]["articles"]      ***REMOVED*** чужое
    assert "Астра ООО" in refs["all"]["counterparties"]

    у_финотдела = client.get("/api/operplan/references?year=2026").json()
    assert "4109_Обучение" in у_финотдела["all"]["articles"]


***REMOVED*** ---------- импорт через API ----------

def test_импорт_листа_через_api(client):
    response = client.post(
        "/api/import/cfo-budgets",
        files={"file": ("бюджеты.xlsx", make_budgets_sheet(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"year": "2026"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["rows"] == 7
    assert client.get("/api/operplan?year=2026").json()["totals"]["rows_count"] == 7


def test_повторный_импорт_заменяет_слой_а_не_дополняет(client):
    payload = {
        "files": {"file": ("бюджеты.xlsx", make_budgets_sheet(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        "data": {"year": "2026"},
    }
    client.post("/api/import/cfo-budgets", **payload)
    second = client.post("/api/import/cfo-budgets", **payload)
    assert second.json()["removed"] == 7
    assert client.get("/api/operplan?year=2026").json()["totals"]["rows_count"] == 7


def test_импорт_доступен_только_финотделу(client):
    cfo_id = repo.ensure_cfo(HR)
    response = client.post(
        "/api/import/cfo-budgets",
        files={"file": ("бюджеты.xlsx", make_budgets_sheet(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"year": "2026"},
        headers={"X-User": f"cfo:{cfo_id}"},
    )
    assert response.status_code == 403


def test_предпросмотр_показывает_расхождения(client):
    wb = openpyxl.load_workbook(io.BytesIO(make_budgets_sheet()))
    wb["Бюджеты ЦФО"].cell(row=9, column=7, value=1)
    buffer = io.BytesIO()
    wb.save(buffer)
    response = client.post(
        "/api/import/inspect",
        files={"file": ("бюджеты.xlsx", buffer.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"kind": "cfo-budgets"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["rows"] == 7
    assert len(body["mismatches"]) == 1


def test_лист_без_нужных_колонок_отклоняется(client):
    wb = openpyxl.Workbook()
    wb.active.cell(row=1, column=1, value="Чужой лист")
    buffer = io.BytesIO()
    wb.save(buffer)
    response = client.post(
        "/api/import/cfo-budgets",
        files={"file": ("чужой.xlsx", buffer.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"year": "2026"},
    )
    assert response.status_code == 422


***REMOVED*** ---------- миграция старой базы ----------

def test_миграция_добавляет_слой_и_не_трогает_данные(tmp_path, monkeypatch):
    """База, созданная до появления слоёв, должна открыться без потерь."""
    import sqlite3

    path = tmp_path / "legacy.db"
    conn = sqlite3.connect(path)
    conn.executescript("""
        CREATE TABLE cfo (id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE,
                          parent_id INTEGER, active INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE budget (id INTEGER PRIMARY KEY, cfo_id INTEGER NOT NULL REFERENCES cfo(id),
                             year INTEGER NOT NULL, name TEXT NOT NULL,
                             UNIQUE (cfo_id, year, name));
        CREATE TABLE budget_row (id INTEGER PRIMARY KEY,
                                 budget_id INTEGER NOT NULL REFERENCES budget(id),
                                 row_no INTEGER, article TEXT NOT NULL,
                                 counterparty TEXT NOT NULL DEFAULT '',
                                 contract TEXT NOT NULL DEFAULT '',
                                 detail TEXT NOT NULL DEFAULT '',
                                 comment TEXT NOT NULL DEFAULT '',
                                 approver_comment TEXT NOT NULL DEFAULT '',
                                 plan_year REAL NOT NULL DEFAULT 0,
                                 carryover INTEGER NOT NULL DEFAULT 1,
                                 source TEXT NOT NULL DEFAULT 'manual',
                                 created_by TEXT NOT NULL DEFAULT '',
                                 created_at TEXT NOT NULL DEFAULT (datetime('now')),
                                 deleted_at TEXT);
        INSERT INTO cfo(id, name) VALUES (1, 'Старое ЦФО');
        INSERT INTO budget(id, cfo_id, year, name) VALUES (7, 1, 2026, 'Основной бюджет');
        INSERT INTO budget_row(id, budget_id, article, plan_year)
            VALUES (42, 7, '4420_ПО', 999000);
    """)
    conn.commit()
    conn.close()

    monkeypatch.setenv("BUDGET_DB", str(path))
    db.set_db_path(path)
    db.init_db()

    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    budget = conn.execute("SELECT * FROM budget WHERE id = 7").fetchone()
    assert budget["layer"] == "main"
    assert budget["name"] == "Основной бюджет"
    row = conn.execute("SELECT * FROM budget_row WHERE id = 42").fetchone()
    assert row["plan_year"] == 999000
    assert row["budget_id"] == 7
    assert conn.execute("PRAGMA foreign_key_check").fetchall() == []
    assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    conn.close()

    ***REMOVED*** повторный запуск безвреден
    db.init_db()
    assert repo.list_rows(2026)[0]["plan_year"] == 999000


***REMOVED*** ---------- Excel: выгрузка, загрузка, факт ----------

def make_fact_sheet(hr_amount: float = 7_000, it_amount: float = 90_000,
                    extra: bool = False) -> bytes:
    """Факт в той же раскладке, что «Бюджеты ЦФО»: суммы — не план, а расход."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Бюджеты ЦФО"
    header = ["Статья оборотов", "ЦФО", "Контрагент", "Договор", "Детализация",
              "Комментарий", "ИТОГО"] + MONTH_NAMES
    for col, value in enumerate(header, start=1):
        ws.cell(row=3, column=col, value=value)

    def put(row, values, months):
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        ws.cell(row=row, column=7, value=sum(months))
        for m, amount in enumerate(months):
            ws.cell(row=row, column=8 + m, value=amount or None)

    put(4, ["ИТОГО", IT, None, "Бюджет обязательных затрат"], [it_amount] + [0] * 11)
    put(5, ["4411_Оборудование", IT, "Астра ООО", "Договор №1", "ИТ инфраструктура",
            "Серверы"], [it_amount] + [0] * 11)
    put(6, ["ИТОГО", HR, None, None], [hr_amount] + [0] * 11)
    put(7, ["4110_Соцрасходы", HR, "Подрядчик ООО", "Договор №3", None, "Подарки"],
        [hr_amount] + [0] * 11)
    if extra:
        put(8, ["4999_Непредвиденное", HR, None, None, None, "Не было в бюджете"],
            [0, 3_000] + [0] * 10)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_выгрузка_бюджета_читается_импортёром_обратно(client, loaded):
    """Круг замкнут: что выгрузили, то и заливается без потери строк и сумм."""
    before = client.get("/api/operplan?year=2026").json()["totals"]
    data = client.get("/api/export/budget.xlsx?year=2026")
    assert data.status_code == 200
    assert data.headers["content-type"].startswith("application/vnd.openxml")

    again = operplan.import_budgets(data.content, "выгрузка.xlsx", 2026, "Тест")
    after = client.get("/api/operplan?year=2026").json()["totals"]
    assert again["rows"] == before["rows_count"]
    assert after["rows_count"] == before["rows_count"]
    assert after["plan_year"] == pytest.approx(before["plan_year"], rel=1e-9)


def test_руководитель_выгружает_только_свой_цфо(client, loaded):
    hr_id = repo.ensure_cfo(HR)
    content = client.get("/api/export/budget.xlsx?year=2026",
                         headers={"X-User": f"cfo:{hr_id}"}).content
    parsed = cfo_budgets.parse(content)
    assert {row["cfo"] for row in parsed["rows"]} == {HR}


def test_загрузка_факта_ложится_на_строки_плана(client, loaded):
    result = operplan.import_fact(make_fact_sheet(), "факт.xlsx", 2026, "Тест")
    assert result["matched"] == 2 and result["created"] == 0

    kpi = client.get("/api/operplan").json()["kpi"]
    assert kpi["fact_loaded"] is True
    assert kpi["base_fact"] == pytest.approx(97_000)


def test_факт_без_строки_в_плане_заводит_её_и_не_задваивает(client, loaded):
    first = operplan.import_fact(make_fact_sheet(extra=True), "факт.xlsx", 2026, "Тест")
    assert first["created"] == 1
    rows_after_first = len(repo.list_rows(2026, layer=operplan.LAYER))

    second = operplan.import_fact(make_fact_sheet(extra=True), "факт.xlsx", 2026, "Тест")
    assert second["created"] == 0 and second["matched"] == 3
    assert len(repo.list_rows(2026, layer=operplan.LAYER)) == rows_after_first


def test_повтор_факта_переносит_суммы_факта_а_не_бюджета(client, loaded):
    operplan.import_fact(make_fact_sheet(), "факт.xlsx", 2026, "Тест")
    result = client.post("/api/operplan/copy-year",
                         json={"from_year": 2026, "to_year": 2027, "source": "fact"}).json()
    assert result["copied"] == 2          ***REMOVED*** перенеслись только строки с фактом
    plan = client.get("/api/operplan?year=2027").json()
    assert plan["totals"]["plan_year"] == pytest.approx(97_000)


def test_повтор_факта_без_факта_отклоняется(client, loaded):
    response = client.post("/api/operplan/copy-year",
                           json={"from_year": 2026, "to_year": 2027, "source": "fact"})
    assert response.status_code == 422
    assert "не загружен" in response.json()["detail"]


def test_руководитель_повторяет_бюджет_только_своего_цфо(client, loaded):
    hr_id = repo.ensure_cfo(HR)
    response = client.post("/api/operplan/copy-year", json={"from_year": 2026, "to_year": 2027},
                           headers={"X-User": f"cfo:{hr_id}"})
    assert response.status_code == 200

    plan = repo.list_rows(2027, layer=operplan.LAYER)
    assert plan, "строки должны перенестись"
    assert {r["cfo_name"] for r in plan} == {HR}, "чужие ЦФО переноситься не должны"


def test_руководитель_заливает_бюджет_не_стирая_чужие_цфо(client, loaded):
    hr_id = repo.ensure_cfo(HR)
    it_before = [r for r in repo.list_rows(2026, layer=operplan.LAYER)
                 if r["cfo_name"] == IT]

    response = client.post(
        "/api/operplan/import",
        data={"year": "2026", "kind": "plan"},
        files={"file": ("бюджет.xlsx", make_budgets_sheet(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers={"X-User": f"cfo:{hr_id}"},
    )
    assert response.status_code == 200
    assert response.json()["foreign"] == 3       ***REMOVED*** три строки ИТ в файле пропущены

    it_after = [r for r in repo.list_rows(2026, layer=operplan.LAYER)
                if r["cfo_name"] == IT]
    assert len(it_after) == len(it_before), "чужие строки не должны стираться"


***REMOVED*** ---------- шаблон Excel ----------

def test_шаблон_отдаёт_лист_бюджеты_цфо_с_блоками_и_справочником(client, loaded):
    response = client.get("/api/export/budget-template.xlsx?year=2027")
    assert response.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Бюджеты ЦФО", "Как заполнять", "Справочник"]

    ws = wb["Бюджеты ЦФО"]
    header = [ws.cell(row=3, column=c).value for c in range(1, 9)]
    assert header == ["Статья оборотов", "ЦФО", "Классификатор", "Контрагент",
                      "Договор", "Детализация", "Комментарий", "ИТОГО"]
    ***REMOVED*** блоки ЦФО есть, строки под ними пустые
    blocks = [ws.cell(row=r, column=2).value for r in range(4, ws.max_row + 1)
              if ws.cell(row=r, column=1).value == "ИТОГО"]
    assert IT in blocks and HR in blocks
    assert not any(ws.cell(row=r, column=7).value for r in range(4, ws.max_row + 1))


def test_шаблон_руководителю_только_его_цфо(client, loaded):
    hr_id = repo.ensure_cfo(HR)
    response = client.get("/api/export/budget-template.xlsx?year=2027",
                          headers={"X-User": f"cfo:{hr_id}"})
    assert response.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(response.content))["Бюджеты ЦФО"]
    blocks = {ws.cell(row=r, column=2).value for r in range(4, ws.max_row + 1)
              if ws.cell(row=r, column=1).value == "ИТОГО"}
    assert blocks == {HR}


def test_заполненный_шаблон_загружается_обратно(client, loaded):
    """Шаблон → заполнили две строки → загрузили: план собрался из них."""
    template = client.get("/api/export/budget-template.xlsx?year=2027").content
    wb = openpyxl.load_workbook(io.BytesIO(template))
    ws = wb["Бюджеты ЦФО"]
    start = next(r for r in range(4, ws.max_row + 1)
                 if ws.cell(row=r, column=1).value == "ИТОГО"
                 and ws.cell(row=r, column=2).value == HR) + 1
    ws.cell(row=start, column=1, value="4109_Обучение")
    ws.cell(row=start, column=7, value=60_000)
    ws.cell(row=start, column=8, value=60_000)
    ***REMOVED*** вторая строка: ИТОГО оставлено пустым, заполнены только месяцы
    ws.cell(row=start + 1, column=1, value="4110_Соцрасходы")
    ws.cell(row=start + 1, column=9, value=25_000)
    ws.cell(row=start + 1, column=10, value=15_000)
    buffer = io.BytesIO()
    wb.save(buffer)

    response = client.post(
        "/api/operplan/import",
        data={"year": "2027", "kind": "plan"},
        files={"file": ("шаблон.xlsx", buffer.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200
    assert response.json()["rows"] == 2, "пустые заготовки не должны попасть в план"

    plan = client.get("/api/operplan?year=2027").json()
    assert plan["totals"]["plan_year"] == pytest.approx(100_000)
    rows = {r["article"]: r for b in plan["blocks"] for r in b["rows"]}
    ***REMOVED*** ИТОГО пустой строки собрано из месяцев, а не осталось нулём
    assert rows["4110_Соцрасходы"]["plan_year"] == pytest.approx(40_000)


***REMOVED*** ---------- область замены при загрузке ----------

def make_one_cfo_sheet(cfo: str, article: str, amount: float) -> bytes:
    """Лист «Бюджеты ЦФО» с одной дирекцией и одной строкой."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Бюджеты ЦФО"
    header = ["Статья оборотов", "ЦФО", "Контрагент", "Договор", "Детализация",
              "Комментарий", "ИТОГО"] + MONTH_NAMES
    for col, value in enumerate(header, start=1):
        ws.cell(row=3, column=col, value=value)
    ws.cell(row=4, column=1, value="ИТОГО")
    ws.cell(row=4, column=2, value=cfo)
    ws.cell(row=4, column=7, value=amount)
    ws.cell(row=5, column=1, value=article)
    ws.cell(row=5, column=2, value=cfo)
    ws.cell(row=5, column=7, value=amount)
    ws.cell(row=5, column=8, value=amount)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_финотдел_заливая_один_цфо_не_стирает_остальные(client, loaded):
    """Файл одной дирекции заменяет только её план, даже у финотдела."""
    it_before = [r for r in repo.list_rows(2026, layer=operplan.LAYER)
                 if r["cfo_name"] == IT]
    assert it_before

    result = client.post(
        "/api/operplan/import",
        data={"year": "2026", "kind": "plan"},
        files={"file": ("дуп.xlsx", make_one_cfo_sheet(HR, "4109_Обучение", 60_000),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    assert result["replaced_cfo"] == 1

    rows = repo.list_rows(2026, layer=operplan.LAYER)
    assert len([r for r in rows if r["cfo_name"] == IT]) == len(it_before), \
        "ИТ в файле не было — его план трогать нельзя"
    hr_rows = [r for r in rows if r["cfo_name"] == HR]
    assert len(hr_rows) == 1 and hr_rows[0]["article"] == "4109_Обучение"


def test_пустой_файл_не_стирает_план(client, loaded):
    """Загрузка шаблона без заполненных строк — не способ снести год."""
    before = len(repo.list_rows(2026, layer=operplan.LAYER))
    template = client.get("/api/export/budget-template.xlsx?year=2026").content
    result = client.post(
        "/api/operplan/import",
        data={"year": "2026", "kind": "plan"},
        files={"file": ("шаблон.xlsx", template,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    assert result["rows"] == 0 and result["removed"] == 0
    assert len(repo.list_rows(2026, layer=operplan.LAYER)) == before


def test_факт_по_одному_цфо_не_гасит_факт_остальных(client, loaded):
    """Загрузка факта по одной дирекции не должна обнулять факт остальных."""
    def fact_by_cfo() -> dict[str, float]:
        months = repo.row_fact_months(2026, layer=operplan.LAYER)
        out: dict[str, float] = {}
        for row in repo.list_rows(2026, layer=operplan.LAYER):
            amount = sum(months.get(int(row["id"]), []))
            if amount:
                out[row["cfo_name"]] = out.get(row["cfo_name"], 0) + amount
        return out

    operplan.import_fact(make_fact_sheet(), "факт.xlsx", 2026, "Тест")
    before = fact_by_cfo()
    assert IT in before and HR in before

    operplan.import_fact(make_one_cfo_sheet(HR, "4109_Обучение", 60_000),
                         "факт-дуп.xlsx", 2026, "Тест", cfo_ids=[repo.ensure_cfo(HR)])
    after = fact_by_cfo()
    assert after[IT] == pytest.approx(before[IT]), "факт ИТ трогать нельзя"
    assert after[HR] == pytest.approx(60_000), "факт ДУП заменён целиком"


def test_повторная_загрузка_факта_убирает_суммы_выпавших_строк(client, loaded):
    """Строка выпала из файла факта — прежний факт на ней остаться не должен."""
    operplan.import_fact(make_fact_sheet(), "факт.xlsx", 2026, "Тест")
    hr_id = repo.ensure_cfo(HR)
    operplan.import_fact(make_one_cfo_sheet(HR, "4109_Обучение", 60_000),
                         "факт-дуп.xlsx", 2026, "Тест", cfo_ids=[hr_id])

    months = repo.row_fact_months(2026, layer=operplan.LAYER)
    stale = [r for r in repo.list_rows(2026, layer=operplan.LAYER)
             if r["cfo_name"] == HR and r["article"] != "4109_Обучение"
             and sum(months.get(int(r["id"]), []))]
    assert not stale, "факт остался на строках, которых в новом файле нет"


***REMOVED*** ---------- пустая строка, заводимая кнопкой в заголовке блока ----------

def test_строка_заводится_без_статьи(client, loaded):
    """Кнопка «+ расход» создаёт пустую строку — заполняют её уже в таблице."""
    hr_id = repo.ensure_cfo(HR)
    response = client.post("/api/operplan/rows?year=2027",
                           json={"cfo_id": hr_id, "budget_name": "Основной бюджет",
                                 "article": ""})
    assert response.status_code == 201
    row_id = response.json()["id"]

    row = repo.get_row(row_id)
    assert row["article"] == "" and row["plan_year"] == 0
    assert row["cfo_name"] == HR


def test_пустая_строка_дозаполняется_правкой(client, loaded):
    """Статья, контрагент и месяцы дописываются в уже созданную строку."""
    hr_id = repo.ensure_cfo(HR)
    row_id = client.post("/api/operplan/rows?year=2027",
                         json={"cfo_id": hr_id, "article": ""}).json()["id"]

    assert client.patch(f"/api/rows/{row_id}",
                        json={"article": "4109_Обучение"}).status_code == 200
    assert client.patch(f"/api/rows/{row_id}",
                        json={"counterparty": "Подрядчик ООО"}).status_code == 200
    ***REMOVED*** равномерная раскладка 120 000 по месяцам: ИТОГО пересчитывается по ним
    months = [10_000.0] * 12
    result = client.patch(f"/api/rows/{row_id}", json={"months": months}).json()
    assert result["row"]["plan_year"] == pytest.approx(120_000)
    assert result["row"]["article"] == "4109_Обучение"


def test_правка_не_даёт_стереть_статью(client, planned):
    """Пустую строку завести можно, а обнулить статью у заполненной — нет."""
    response = client.patch(f"/api/rows/{plan_row()['id']}", json={"article": "  "})
    assert response.status_code == 422


def test_руководитель_не_заводит_строку_в_чужом_цфо(client, loaded):
    hr_id = repo.ensure_cfo(HR)
    it_id = repo.ensure_cfo(IT)
    response = client.post("/api/operplan/rows?year=2027",
                           json={"cfo_id": it_id, "article": ""},
                           headers={"X-User": f"cfo:{hr_id}"})
    assert response.status_code == 403


***REMOVED*** ---------- нумерация строк ----------

def test_сохранить_закрепляет_нумерацию_по_порядку_таблицы(client, planned):
    """До «Сохранить» номера тянутся из файла; кнопка проставляет их подряд."""
    result = client.post("/api/operplan/renumber", json={"year": 2027}).json()
    assert result["renumbered"] > 0

    after = repo.list_rows(2027, layer=operplan.LAYER)
    by_budget: dict[int, list[dict]] = {}
    for row in after:
        by_budget.setdefault(int(row["budget_id"]), []).append(row)

    for rows in by_budget.values():
        ordered = sorted(rows, key=lambda r: int(r["row_no"]))
        assert [int(r["row_no"]) for r in ordered] == list(range(1, len(rows) + 1)), \
            "внутри бюджета номера должны идти подряд с единицы"
        keys = [((r["article"] or "").lower(), (r["counterparty"] or "").lower())
                for r in ordered]
        assert keys == sorted(keys), "и совпадать с порядком таблицы: статья, контрагент"


def test_повторное_сохранение_ничего_не_переставляет(client, planned):
    client.post("/api/operplan/renumber", json={"year": 2027})
    numbers = {int(r["id"]): int(r["row_no"])
               for r in repo.list_rows(2027, layer=operplan.LAYER)}

    again = client.post("/api/operplan/renumber", json={"year": 2027}).json()
    assert again["renumbered"] == 0, "порядок уже закреплён — менять нечего"
    assert {int(r["id"]): int(r["row_no"])
            for r in repo.list_rows(2027, layer=operplan.LAYER)} == numbers


def test_замок_нумерацию_не_трогает(client, planned):
    numbers = {int(r["id"]): r["row_no"]
               for r in repo.list_rows(2027, layer=operplan.LAYER)}
    result = client.post("/api/operplan/lock", json={"year": 2027, "locked": True}).json()
    assert "renumbered" not in result
    assert {int(r["id"]): r["row_no"]
            for r in repo.list_rows(2027, layer=operplan.LAYER)} == numbers


def test_руководитель_нумерует_только_свои_блоки(client, planned):
    hr_id = repo.ensure_cfo(HR)
    it_before = {int(r["id"]): r["row_no"]
                 for r in repo.list_rows(2027, layer=operplan.LAYER)
                 if r["cfo_name"] == IT}

    response = client.post("/api/operplan/renumber", json={"year": 2027},
                           headers={"X-User": f"cfo:{hr_id}"})
    assert response.status_code == 200

    assert {int(r["id"]): r["row_no"]
            for r in repo.list_rows(2027, layer=operplan.LAYER)
            if r["cfo_name"] == IT} == it_before, "чужие номера трогать нельзя"
    hr_numbers = sorted(int(r["row_no"]) for r in repo.list_rows(2027, layer=operplan.LAYER)
                        if r["cfo_name"] == HR)
    assert hr_numbers == list(range(1, len(hr_numbers) + 1))


def test_нумерация_прошлого_года_отклоняется(client, planned):
    response = client.post("/api/operplan/renumber", json={"year": 2026})
    assert response.status_code == 403
    assert "база для сравнения" in response.json()["detail"]


def test_перенумерация_не_трогает_другой_год_и_слой(client, planned):
    base_before = {int(r["id"]): r["row_no"]
                   for r in repo.list_rows(2026, layer=operplan.LAYER)}
    main_before = {int(r["id"]): r["row_no"] for r in repo.list_rows(2026)}

    client.post("/api/operplan/renumber", json={"year": 2027})

    assert {int(r["id"]): r["row_no"]
            for r in repo.list_rows(2026, layer=operplan.LAYER)} == base_before
    assert {int(r["id"]): r["row_no"] for r in repo.list_rows(2026)} == main_before


def test_пустая_строка_при_перенумерации_идёт_первой(client, planned):
    """Незаполненную строку таблица держит сверху блока — номер должен совпасть."""
    it_id = repo.ensure_cfo(IT)
    row_id = client.post("/api/operplan/rows?year=2027",
                         json={"cfo_id": it_id, "budget_name": "Бюджет развития",
                               "article": ""}).json()["id"]

    client.post("/api/operplan/renumber", json={"year": 2027})
    assert int(repo.get_row(row_id)["row_no"]) == 1


***REMOVED*** ---------- круг: строка с нулевой суммой ----------

def test_нулевая_строка_со_статьёй_переживает_круг(client, loaded):
    """Расход заведён, сумму ещё не проставили — выгрузка и заливка её не теряют.

    Правило «нет суммы — строка техническая» писалось под исходный файл
    ликвидности, где такие строки действительно пустые. На круговой заливке
    собственной выгрузки оно съедало законную строку плана.
    """
    cfo_id = repo.ensure_cfo(HR)
    budget_id = repo.ensure_budget(cfo_id, 2027, "Основной бюджет", layer=operplan.LAYER)
    repo.create_row(budget_id, {
        "article": "4505_Выплата доли девелопера", "counterparty": "Сигма ООО",
        "contract": "", "detail": "", "comment": "", "approver_comment": "",
        "plan_year": 0.0,
    }, months=[0.0] * 12, author="Тест", source="manual", row_no=1)

    before = client.get("/api/operplan?year=2027").json()["totals"]["rows_count"]
    content = client.get("/api/export/budget.xlsx?year=2027").content
    result = operplan.import_budgets(content, "выгрузка.xlsx", 2027, "Тест")

    assert result["rows"] == before, "нулевая строка не должна пропасть при заливке"
    rows = repo.list_rows(2027, layer=operplan.LAYER)
    zero = [r for r in rows if r["article"] == "4505_Выплата доли девелопера"]
    assert len(zero) == 1
    assert zero[0]["counterparty"] == "Сигма ООО"
    assert zero[0]["plan_year"] == 0


def test_строка_без_статьи_круг_не_переживает(client, loaded):
    """Заготовка от кнопки «+ расход» — заливать нечего, и она отсеивается.

    Уточнение, которое стоит помнить: круг замены — дирекции, чьи строки есть
    в файле. Заготовка исчезает потому, что её блок перезаписывается соседней
    заполненной строкой. Блок, где заготовка одна, в замену не попадает вовсе
    (см. test_одинокая_заготовка_остаётся_на_месте).
    """
    cfo_id = repo.ensure_cfo(HR)
    budget_id = repo.ensure_budget(cfo_id, 2027, "Основной бюджет", layer=operplan.LAYER)
    for row_no, article in ((1, ""), (2, "4109_Обучение")):
        repo.create_row(budget_id, {
            "article": article, "counterparty": "", "contract": "", "detail": "",
            "comment": "", "approver_comment": "", "plan_year": 0.0 if not article else 5000.0,
        }, months=[0.0] * 12, author="Тест", source="manual", row_no=row_no)

    content = client.get("/api/export/budget.xlsx?year=2027").content
    result = operplan.import_budgets(content, "выгрузка.xlsx", 2027, "Тест")
    assert result["rows"] == 1
    articles = [r["article"] for r in repo.list_rows(2027, layer=operplan.LAYER)]
    assert articles == ["4109_Обучение"], "заготовка без статьи залиться не должна"


def test_одинокая_заготовка_остаётся_на_месте(client, loaded):
    """В блоке только заготовка — заливать в него нечего, и он не трогается.

    Это не потеря и не дубль: замена идёт по дирекциям, чьи строки есть в файле,
    а у этой их нет. Иначе выгрузка одной дирекции стирала бы чужие бюджеты.
    """
    cfo_id = repo.ensure_cfo(HR)
    budget_id = repo.ensure_budget(cfo_id, 2027, "Основной бюджет", layer=operplan.LAYER)
    repo.create_row(budget_id, {
        "article": "", "counterparty": "", "contract": "", "detail": "",
        "comment": "", "approver_comment": "", "plan_year": 0.0,
    }, months=[0.0] * 12, author="Тест", source="manual", row_no=1)

    content = client.get("/api/export/budget.xlsx?year=2027").content
    result = operplan.import_budgets(content, "выгрузка.xlsx", 2027, "Тест")
    assert result["rows"] == 0
    assert len(repo.list_rows(2027, layer=operplan.LAYER)) == 1


def test_факт_нулевые_строки_не_заводит(client, loaded):
    """В факте ноль — это не расход: пустую строку в плане он заводить не должен."""
    wb = openpyxl.load_workbook(io.BytesIO(make_fact_sheet()))
    ws = wb["Бюджеты ЦФО"]
    last = ws.max_row + 1
    ws.cell(row=last, column=1, value="4999_Строка без расхода")
    ws.cell(row=last, column=2, value=HR)
    ws.cell(row=last, column=3, value="Никто ООО")
    ws.cell(row=last, column=7, value=0)
    buffer = io.BytesIO()
    wb.save(buffer)

    before = len(repo.list_rows(2026, layer=operplan.LAYER))
    result = operplan.import_fact(buffer.getvalue(), "факт.xlsx", 2026, "Тест")
    assert result["created"] == 0
    assert len(repo.list_rows(2026, layer=operplan.LAYER)) == before


***REMOVED*** ---------- классификатор МРГ/МРК ----------

def test_новая_строка_по_умолчанию_мрг(client, loaded):
    """Выбирать за пользователя нечего — строка заводится как МРГ."""
    cfo_id = repo.ensure_cfo(HR)
    response = client.post("/api/operplan/rows?year=2027", json={
        "cfo_id": cfo_id, "article": "4109_Обучение", "plan_year": 1000,
    })
    assert response.status_code == 201, response.text
    assert repo.get_row(response.json()["id"])["classifier"] == "МРГ"


def test_классификатор_правится_и_знает_только_два_значения(client, planned):
    row = plan_row()
    assert client.patch(f"/api/rows/{row['id']}", json={"classifier": "МРК"}).status_code == 200
    assert repo.get_row(row["id"])["classifier"] == "МРК"

    bad = client.patch(f"/api/rows/{row['id']}", json={"classifier": "МРФ"})
    assert bad.status_code == 422
    assert repo.get_row(row["id"])["classifier"] == "МРК", "отказ не должен ничего менять"


def test_классификатор_виден_на_вкладке(client, planned):
    row = plan_row()
    client.patch(f"/api/rows/{row['id']}", json={"classifier": "МРК"})
    rows = {r["id"]: r for b in client.get("/api/operplan?year=2027").json()["blocks"]
            for r in b["rows"]}
    assert rows[row["id"]]["classifier"] == "МРК"
    assert all(r["classifier"] in ("МРГ", "МРК") for r in rows.values())


def test_повтор_года_переносит_классификатор(client, loaded):
    base = repo.list_rows(2026, layer=operplan.LAYER)[0]
    repo.update_row(base["id"], {"classifier": "МРК"})
    operplan.copy_year("Тест", 2026, 2027)

    copied = [r for r in repo.list_rows(2027, layer=operplan.LAYER)
              if r["article"] == base["article"] and r["counterparty"] == base["counterparty"]]
    assert copied and copied[0]["classifier"] == "МРК"


def test_выгрузка_несёт_классификатор_и_круг_его_не_теряет(client, planned):
    row = plan_row()
    client.patch(f"/api/rows/{row['id']}", json={"classifier": "МРК"})

    content = client.get("/api/export/budget.xlsx?year=2027").content
    ws = openpyxl.load_workbook(io.BytesIO(content))["Бюджеты ЦФО"]
    assert ws.cell(row=3, column=3).value == "Классификатор"
    ***REMOVED*** ключ по статье и контрагенту: одна статья в плане встречается не раз
    values = {(ws.cell(row=r, column=1).value, ws.cell(row=r, column=4).value):
              ws.cell(row=r, column=3).value
              for r in range(4, ws.max_row + 1)
              if ws.cell(row=r, column=1).value not in (None, "ИТОГО")}
    assert values[(row["article"], row["counterparty"] or None)] == "МРК"
    assert set(values.values()) <= {"МРГ", "МРК"}

    operplan.import_budgets(content, "выгрузка.xlsx", 2027, "Тест")
    again = [r for r in repo.list_rows(2027, layer=operplan.LAYER)
             if r["article"] == row["article"] and r["counterparty"] == row["counterparty"]]
    assert again and again[0]["classifier"] == "МРК", "круг не должен терять классификатор"


def test_загрузка_чинит_чужое_значение_и_переживает_отсутствие_колонки(client, loaded):
    """Файл из чужих рук важнее описки в одной колонке — строка не теряется."""
    wb = openpyxl.load_workbook(io.BytesIO(make_budgets_sheet()))
    ws = wb["Бюджеты ЦФО"]
    ***REMOVED*** у фикстуры колонки нет вовсе — сначала грузим как есть
    operplan.import_budgets(make_budgets_sheet(), "без колонки.xlsx", 2027, "Тест")
    assert {r["classifier"] for r in repo.list_rows(2027, layer=operplan.LAYER)} == {"МРГ"}

    ***REMOVED*** теперь вставляем колонку и пишем в неё мусор и строчные буквы
    ws.insert_cols(3)
    ws.cell(row=3, column=3, value="Классификатор")
    filled = [r for r in range(4, ws.max_row + 1)
              if ws.cell(row=r, column=1).value not in (None, "ИТОГО")]
    ws.cell(row=filled[0], column=3, value="мрк")
    ws.cell(row=filled[1], column=3, value="Чужое ЮЛ")
    buffer = io.BytesIO()
    wb.save(buffer)

    before = len(repo.list_rows(2027, layer=operplan.LAYER))
    operplan.import_budgets(buffer.getvalue(), "с колонкой.xlsx", 2027, "Тест")
    rows = repo.list_rows(2027, layer=operplan.LAYER)
    assert len(rows) == before, "строка с непонятным значением не должна пропасть"
    assert sorted({r["classifier"] for r in rows}) == ["МРГ", "МРК"]


def test_шаблон_даёт_выбрать_из_двух(client, loaded):
    content = client.get("/api/export/budget-template.xlsx?year=2027").content
    ws = openpyxl.load_workbook(io.BytesIO(content))["Бюджеты ЦФО"]
    rules = [r for r in ws.data_validations.dataValidation
             if str(r.sqref).startswith("C")]
    assert len(rules) == 1, "на колонке классификатора один список"
    assert rules[0].formula1 == '"МРГ,МРК"'
    assert str(rules[0].sqref).startswith("C4")


def _range_values(wb, name: str) -> list[str]:
    """Что реально лежит в именованном диапазоне, на который смотрит список."""
    ref = str(wb.defined_names[name].value)
    sheet, _, cells = ref.rpartition("!")
    ws = wb[sheet.strip("'")]
    return [c.value for row in ws[cells.replace("$", "")] for c in row]


def test_шаблон_подсказывает_статьи_и_контрагентов_из_справочника(client, loaded):
    """У каждого блока ЦФО свой список — из того, что уже встречалось у дирекции."""
    content = client.get("/api/export/budget-template.xlsx?year=2027").content
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["Бюджеты ЦФО"]
    ***REMOVED*** строка, с которой начинается блок ДУП
    start = next(r for r in range(4, ws.max_row + 1)
                 if ws.cell(row=r, column=1).value == "ИТОГО"
                 and ws.cell(row=r, column=2).value == HR) + 1

    def rule_at(column: str):
        return next(r for r in ws.data_validations.dataValidation
                    if f"{column}{start}" in r.sqref)

    articles = rule_at("A")
    assert articles.type == "list"
    assert not articles.showErrorMessage, "список подсказывает, новую статью не запрещает"
    assert "4109_Обучение" in _range_values(wb, articles.formula1)
    assert "4411_Оборудование" not in _range_values(wb, articles.formula1),         "в блоке ДУП не должно быть статей ИТ"
    assert _range_values(wb, rule_at("D").formula1), "у контрагентов список тоже есть"


def test_шаблон_руководителю_не_подсказывает_чужие_статьи(client, loaded):
    hr_id = repo.ensure_cfo(HR)
    content = client.get("/api/export/budget-template.xlsx?year=2027",
                         headers={"X-User": f"cfo:{hr_id}"}).content
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["Бюджеты ЦФО"]
    everything = {v for name in wb.defined_names for v in _range_values(wb, name)}
    assert "4109_Обучение" in everything
    assert "4411_Оборудование" not in everything, "статьи ИТ руководителю ДУП не видны"
    assert [r for r in ws.data_validations.dataValidation
            if str(r.sqref).startswith("A")], "список статей у блока есть"


def test_миграция_проставляет_мрг_существующим_строкам(tmp_path, monkeypatch):
    """Рабочая база колонки не знает: ALTER доливает её со значением по умолчанию."""
    import sqlite3

    path = tmp_path / "legacy.db"
    conn = sqlite3.connect(path)
    conn.executescript(
        """
        CREATE TABLE cfo (id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE,
                          parent_id INTEGER, active INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE budget (id INTEGER PRIMARY KEY, cfo_id INTEGER NOT NULL REFERENCES cfo(id),
                             year INTEGER NOT NULL, name TEXT NOT NULL, layer TEXT NOT NULL,
                             UNIQUE (cfo_id, year, name, layer));
        CREATE TABLE budget_row (id INTEGER PRIMARY KEY,
                                 budget_id INTEGER NOT NULL REFERENCES budget(id),
                                 row_no INTEGER, article TEXT NOT NULL,
                                 counterparty TEXT NOT NULL DEFAULT '',
                                 contract TEXT NOT NULL DEFAULT '',
                                 detail TEXT NOT NULL DEFAULT '',
                                 comment TEXT NOT NULL DEFAULT '',
                                 approver_comment TEXT NOT NULL DEFAULT '',
                                 plan_year REAL NOT NULL DEFAULT 0,
                                 carryover INTEGER NOT NULL DEFAULT 1,
                                 source TEXT NOT NULL DEFAULT 'manual',
                                 created_by TEXT NOT NULL DEFAULT '',
                                 created_at TEXT NOT NULL DEFAULT (datetime('now')),
                                 deleted_at TEXT);
        INSERT INTO cfo(id, name) VALUES (1, 'Старое ЦФО');
        INSERT INTO budget(id, cfo_id, year, name, layer)
            VALUES (7, 1, 2026, 'Основной бюджет', 'operplan');
        INSERT INTO budget_row(id, budget_id, article, plan_year) VALUES (42, 7, '4420_ПО', 999);
        """
    )
    conn.commit()
    conn.close()

    monkeypatch.setenv("BUDGET_DB", str(path))
    db.set_db_path(path)
    db.init_db()

    row = repo.get_row(42)
    assert row["classifier"] == "МРГ"
    assert row["plan_year"] == 999
    db.init_db()                            ***REMOVED*** повторный запуск безвреден
    assert repo.get_row(42)["classifier"] == "МРГ"


def test_отчёт_загрузки_совпадает_с_тем_что_легло(client, loaded):
    """Сумма в ответе считалась по колонке ИТОГО файла, а строки писались
    с суммой по месяцам, когда ИТОГО пуст. Цифры расходились."""
    wb = openpyxl.load_workbook(io.BytesIO(make_budgets_sheet()))
    ws = wb["Бюджеты ЦФО"]
    filled = [r for r in range(4, ws.max_row + 1)
              if ws.cell(row=r, column=1).value not in (None, "ИТОГО")]
    ***REMOVED*** У первой строки стираем ИТОГО и оставляем один месяц. Чистится присваиванием
    ***REMOVED*** в .value: openpyxl в ws.cell(..., value=None) понимает None как «не задавать»
    ***REMOVED*** и оставляет прежнее значение — на этом тест уже один раз оказался пустым.
    ws.cell(row=filled[0], column=7).value = None
    for c in range(8, 20):
        ws.cell(row=filled[0], column=c).value = None
    ws.cell(row=filled[0], column=8, value=12_000)
    buffer = io.BytesIO()
    wb.save(buffer)

    result = operplan.import_budgets(buffer.getvalue(), "файл.xlsx", 2027, "Тест")
    in_db = sum(r["plan_year"] for r in repo.list_rows(2027, layer=operplan.LAYER))
    assert result["total"] == pytest.approx(in_db)
    assert result["total"] > 0


def test_отчёт_загрузки_не_считает_чужие_цфо(client, loaded):
    """Руководителю чужие строки пропускаются — в сумму они попадать не должны."""
    hr_id = repo.ensure_cfo(HR)
    result = operplan.import_budgets(make_budgets_sheet(), "файл.xlsx", 2027, "Тест",
                                     cfo_ids=[hr_id])
    in_db = sum(r["plan_year"] for r in repo.list_rows(2027, layer=operplan.LAYER))
    assert result["foreign"] > 0, "в файле есть чужие дирекции"
    assert result["total"] == pytest.approx(in_db)
