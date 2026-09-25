"""Тесты API: сквозные сценарии на изолированной базе."""
from __future__ import annotations

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import db, repo, service
from app.main import app


@pytest.fixture()
def client(tmp_path, monkeypatch):
    """Клиент со свежей базой на каждый тест."""
    monkeypatch.setenv("BUDGET_DB", str(tmp_path / "test.db"))
    db.set_db_path(tmp_path / "test.db")
    db.init_db()
    db.set_setting("year", "2026")
    db.set_setting("closed_through", "7")
    db.set_setting("as_of", "2026-08-04")
    with TestClient(app) as c:
        yield c


@pytest.fixture()
def budget(client):
    """ЦФО с одной строкой бюджета: 1 200 000 ₽, равномерно по месяцам."""
    cfo_id = repo.ensure_cfo("Первая дирекция")
    response = client.post("/api/rows", json={
        "cfo_id": cfo_id, "budget_name": "Бюджет обязательных затрат",
        "article": "4420_Расходы на приобретение ПО", "counterparty": "Астра ООО",
        "plan_year": 1_200_000, "spread": "even",
    })
    assert response.status_code == 201, response.text
    return {"cfo_id": cfo_id, "row_id": response.json()["id"]}


***REMOVED*** ---------- базовое ----------

def test_health_и_настройки(client):
    assert client.get("/health").json()["status"] == "ok"
    cfg = client.get("/api/settings").json()
    assert cfg["year"] == 2026
    assert cfg["closed_through"] == 7
    assert 0 < cfg["pace"] < 1


def test_по_умолчанию_роль_финотдела(client):
    me = client.get("/api/me").json()
    assert me["is_fin"] is True


***REMOVED*** ---------- строки бюджета ----------

def test_создание_строки_разбивает_план_по_месяцам(client, budget):
    card = client.get(f"/api/cfo/{budget['cfo_id']}").json()
    row = card["budgets"][0]["rows"][0]
    assert row["plan_year"] == 1_200_000
    assert row["plan_months"][0] == pytest.approx(100_000)
    assert sum(row["plan_months"]) == pytest.approx(1_200_000)


def test_строка_с_пустой_статьёй_отклоняется(client, budget):
    response = client.post("/api/rows", json={
        "cfo_id": budget["cfo_id"], "article": "   ", "plan_year": 100,
    })
    assert response.status_code == 422


def test_удаление_строки_убирает_её_из_карточки_и_пишет_в_журнал(client, budget):
    assert client.delete(f"/api/rows/{budget['row_id']}").status_code == 200
    card = client.get(f"/api/cfo/{budget['cfo_id']}").json()
    assert card["budgets"][0]["rows"] == []
    actions = [a["action"] for a in client.get("/api/audit").json()]
    assert "Удаление строки" in actions


***REMOVED*** ---------- факт и заявки ----------

def test_факт_уменьшает_остаток_по_кассе(client, budget):
    client.post("/api/facts", json={
        "row_id": budget["row_id"], "pay_date": "2026-03-15", "amount": 300_000,
    })
    card = client.get(f"/api/cfo/{budget['cfo_id']}").json()
    row = card["budgets"][0]["rows"][0]
    assert row["fact"] == 300_000
    assert row["cash_remain"] == 900_000
    assert row["fact_months"][2] == 300_000


def test_заявка_уменьшает_свободный_лимит_но_не_факт(client, budget):
    client.post("/api/claims", json={
        "row_id": budget["row_id"], "cfo_id": budget["cfo_id"],
        "pay_date": "2026-09-15", "amount": 500_000, "counterparty": "Астра ООО",
    })
    card = client.get(f"/api/cfo/{budget['cfo_id']}").json()
    row = card["budgets"][0]["rows"][0]
    assert row["fact"] == 0
    assert row["claims"] == 500_000
    assert row["operplan"] == 500_000
    assert row["free_limit"] == 700_000
    assert row["cash_remain"] == 1_200_000


def test_заявка_сверх_лимита_сохраняется_с_предупреждением(client, budget):
    response = client.post("/api/claims", json={
        "row_id": budget["row_id"], "cfo_id": budget["cfo_id"],
        "pay_date": "2026-09-15", "amount": 1_500_000,
    })
    assert response.status_code == 201
    body = response.json()
    assert body["over_limit"] is True
    assert "сверх лимита" in body["warning"]


def test_проверка_лимита_считает_превышение(client, budget):
    client.post("/api/claims", json={
        "row_id": budget["row_id"], "cfo_id": budget["cfo_id"],
        "pay_date": "2026-09-15", "amount": 1_000_000,
    })
    check = client.get(f"/api/rows/{budget['row_id']}/limit", params={"amount": 300_000}).json()
    assert check["free_limit"] == 200_000
    assert check["over_limit"] is True
    assert check["excess"] == 100_000


def test_расход_вне_бюджета_попадает_в_отдельный_реестр(client, budget):
    client.post("/api/facts", json={
        "cfo_id": budget["cfo_id"], "pay_date": "2026-07-10", "amount": 66_917,
        "counterparty": "ПЛАТФОРМИКС ООО", "purpose": "Срочная доработка",
        "out_of_budget": True,
    })
    items = client.get("/api/out-of-budget").json()
    assert len(items) == 1
    assert items[0]["amount"] == 66_917
    dash = client.get("/api/dashboard").json()
    assert dash["out_of_budget_total"] == 66_917


***REMOVED*** ---------- сопоставление ----------

def test_неразобранный_платёж_попадает_в_кандидаты_по_статье_и_контрагенту(client, budget):
    fact_id = repo.create_fact({
        "year": 2026, "pay_date": "2026-08-01", "amount": 250_000,
        "article": "4420_Расходы на приобретение ПО", "counterparty": "Астра ООО",
        "source": "1c", "external_id": "x1",
    })
    unmatched = client.get("/api/facts/unmatched").json()
    assert len(unmatched) == 1

    candidates = client.get(f"/api/facts/{fact_id}/candidates").json()
    assert candidates
    assert candidates[0]["row_id"] == budget["row_id"]
    assert "статья" in candidates[0]["matched_by"]

    assert client.post(f"/api/facts/{fact_id}/match",
                       json={"row_id": budget["row_id"]}).status_code == 200
    assert client.get("/api/facts/unmatched").json() == []
    card = client.get(f"/api/cfo/{budget['cfo_id']}").json()
    assert card["budgets"][0]["rows"][0]["fact"] == 250_000


def test_автосопоставление_по_номеру_заявки(client, budget):
    client.post("/api/claims", json={
        "row_id": budget["row_id"], "cfo_id": budget["cfo_id"], "number_1c": "ЗП-2026-777",
        "pay_date": "2026-09-01", "amount": 100_000,
    })
    repo.create_fact({
        "year": 2026, "pay_date": "2026-09-02", "amount": 100_000,
        "claim_number": "ЗП-2026-777", "counterparty": "Кто-то ещё",
        "source": "1c", "external_id": "x2",
    })
    assert client.post("/api/facts/auto-match").json()["matched"] == 1


***REMOVED*** ---------- перенос лимита ----------

def test_перенос_лимита_между_месяцами(client, budget):
    response = client.post("/api/adjustments", json={
        "kind": "month", "from_row_id": budget["row_id"],
        "from_month": 11, "to_month": 9, "amount": 50_000,
        "reason": "Перенос платежа по согласованию с поставщиком",
    })
    assert response.status_code == 200
    card = client.get(f"/api/cfo/{budget['cfo_id']}").json()
    plan = card["budgets"][0]["rows"][0]["plan_months"]
    assert plan[10] == pytest.approx(50_000)     ***REMOVED*** ноябрь уменьшился
    assert plan[8] == pytest.approx(150_000)     ***REMOVED*** сентябрь вырос
    assert sum(plan) == pytest.approx(1_200_000)  ***REMOVED*** годовой лимит не изменился


def test_перенос_больше_доступного_отклоняется(client, budget):
    response = client.post("/api/adjustments", json={
        "kind": "month", "from_row_id": budget["row_id"],
        "from_month": 11, "to_month": 9, "amount": 999_999,
        "reason": "Проверка ограничения",
    })
    assert response.status_code == 422
    assert "запланировано" in response.json()["detail"]


def test_перенос_без_обоснования_отклоняется(client, budget):
    response = client.post("/api/adjustments", json={
        "kind": "month", "from_row_id": budget["row_id"],
        "from_month": 11, "to_month": 9, "amount": 1000, "reason": "",
    })
    assert response.status_code == 422


def test_перенос_между_цфо_запрещён(client, budget):
    other_cfo = repo.ensure_cfo("Финансовый отдел")
    other = client.post("/api/rows", json={
        "cfo_id": other_cfo, "article": "4106_Аудит", "plan_year": 500_000,
    }).json()["id"]
    response = client.post("/api/adjustments", json={
        "kind": "row", "from_row_id": budget["row_id"], "to_row_id": other,
        "amount": 100_000, "reason": "Попытка переноса между дирекциями",
    })
    assert response.status_code == 422
    assert "ЦФО" in response.json()["detail"]


***REMOVED*** ---------- роли ----------

def test_руководитель_цфо_видит_только_свой_цфо(client, budget):
    other_cfo = repo.ensure_cfo("Финансовый отдел")
    client.post("/api/rows", json={
        "cfo_id": other_cfo, "article": "4106_Аудит", "plan_year": 500_000,
    })
    as_cfo = {"X-User": f"cfo:{budget['cfo_id']}"}

    dash = client.get("/api/dashboard", headers=as_cfo).json()
    assert {c["cfo_id"] for c in dash["cards"]} == {budget["cfo_id"]}
    assert len(client.get("/api/cfo", headers=as_cfo).json()) == 1


def test_руководитель_цфо_не_меняет_настройки_и_не_сопоставляет(client, budget):
    as_cfo = {"X-User": f"cfo:{budget['cfo_id']}"}
    assert client.put("/api/settings", json={"closed_through": 5}, headers=as_cfo).status_code == 403
    assert client.get("/api/facts/unmatched", headers=as_cfo).status_code == 403


def test_руководитель_цфо_не_добавляет_строки_в_чужой_цфо(client, budget):
    other_cfo = repo.ensure_cfo("Финансовый отдел")
    response = client.post("/api/rows", json={
        "cfo_id": other_cfo, "article": "4106_Аудит", "plan_year": 100_000,
    }, headers={"X-User": f"cfo:{budget['cfo_id']}"})
    assert response.status_code == 403


def test_руководителю_цфо_закрыты_журнал_и_история_загрузок(client, budget):
    """Вкладки «Корректировки» и «Импорт / экспорт» — только финотдел."""
    as_cfo = {"X-User": f"cfo:{budget['cfo_id']}"}
    assert client.get("/api/audit", headers=as_cfo).status_code == 403
    assert client.get("/api/imports", headers=as_cfo).status_code == 403
    assert client.get("/api/audit").status_code == 200
    assert client.get("/api/imports").status_code == 200


def test_руководитель_цфо_не_смотрит_лимит_чужой_строки(client, budget):
    other_cfo = repo.ensure_cfo("Финансовый отдел")
    made = client.post("/api/rows", json={
        "cfo_id": other_cfo, "article": "4106_Аудит", "plan_year": 500_000,
    }).json()
    as_cfo = {"X-User": f"cfo:{budget['cfo_id']}"}
    assert client.get(f"/api/rows/{made['id']}/limit?amount=1", headers=as_cfo).status_code == 403
    own = client.get(f"/api/rows/{budget['row_id']}/limit?amount=1", headers=as_cfo)
    assert own.status_code == 200


***REMOVED*** ---------- дашборд и тепловая карта ----------

def test_дашборд_считает_итоги_и_статусы(client, budget):
    client.post("/api/facts", json={
        "row_id": budget["row_id"], "pay_date": "2026-02-15", "amount": 1_150_000,
    })
    dash = client.get("/api/dashboard").json()
    assert dash["totals"]["plan_year"] == 1_200_000
    assert dash["totals"]["fact"] == 1_150_000
    card = dash["cards"][0]
    assert card["status"] == "warning"      ***REMOVED*** 96% при ожидаемых ~62%
    assert card["totals"]["free_limit"] == 50_000


def test_тепловая_карта_отдаёт_отклонения_только_по_закрытым_месяцам(client, budget):
    client.post("/api/facts", json={
        "row_id": budget["row_id"], "pay_date": "2026-01-20", "amount": 150_000,
    })
    rows = client.get("/api/heatmap").json()["rows"]
    cells = rows[0]["cells"]
    assert cells[0]["closed"] is True
    assert cells[0]["deviation"] == pytest.approx(0.5)   ***REMOVED*** 150k против плана 100k
    assert cells[11]["closed"] is False
    assert cells[11]["deviation"] is None


***REMOVED*** ---------- импорт и экспорт ----------

def _operplan_file(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Оперплан ЦФО"
    months = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
              "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
    for m in range(12):
        ws.cell(row=3, column=16 + m, value="ФАКТ" if m < 4 else "ПРОГНОЗ")
    header = ["№", "Статья оборотов", "ЦФО", "Контрагент", "Договор", "Детализация",
              "Комментарий", "Комм наш", "Бюджет - опер план", "БЮДЖЕТ", "Оперплан фикс",
              "", "Итого опер план", "Факт", "Остаток бюжета"] + months
    for c, value in enumerate(header, start=1):
        ws.cell(row=5, column=c, value=value)
    ws.cell(row=6, column=2, value="ИТОГО")
    ws.cell(row=6, column=3, value="Десятая дирекция")
    ws.cell(row=6, column=5, value="Бюджет обязательных затрат")
    ws.cell(row=6, column=10, value=600_000)
    ws.cell(row=7, column=1, value=1)
    ws.cell(row=7, column=2, value="4301_Анализ рынка")
    ws.cell(row=7, column=3, value="Десятая дирекция")
    ws.cell(row=7, column=10, value=600_000)
    for m, amount in enumerate([100_000, 100_000, 0, 0, 0, 0, 0, 0, 200_000, 0, 0, 200_000]):
        if amount:
            ws.cell(row=7, column=16 + m, value=amount)
    wb.save(path)
    return path


def test_импорт_оперплана_создаёт_строки_факт_и_заявки(client, tmp_path):
    path = _operplan_file(tmp_path / "op.xlsx")
    with open(path, "rb") as fh:
        response = client.post("/api/import/operplan", data={"year": 2026},
                               files={"file": ("op.xlsx", fh, "application/vnd.ms-excel")})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["rows"] == 1
    assert body["closed_through"] == 4
    assert body["facts"] == 2      ***REMOVED*** январь и февраль — закрытые месяцы
    assert body["claims"] == 2     ***REMOVED*** сентябрь и декабрь — прогноз

    dash = client.get("/api/dashboard").json()
    assert dash["totals"]["plan_year"] == 600_000
    assert dash["totals"]["fact"] == 200_000
    assert dash["totals"]["claims"] == 400_000


def test_предпросмотр_файла_дирекции_возвращает_сопоставление(client, tmp_path):
    from tests.test_importers import make_directorate

    path = make_directorate(tmp_path / "d.xlsx")
    with open(path, "rb") as fh:
        response = client.post("/api/import/inspect", data={"kind": "plan"},
                               files={"file": ("d.xlsx", fh, "application/vnd.ms-excel")})
    assert response.status_code == 200
    body = response.json()
    assert body["header_row"] == 3
    assert body["mapping"]["article"] == 0
    assert body["months_found"] == 12


def test_экспорт_оперплана_отдаёт_xlsx(client, budget):
    response = client.get("/api/export/operplan.xlsx")
    assert response.status_code == 200
    assert response.content[:2] == b"PK"          ***REMOVED*** zip-контейнер xlsx
    assert "attachment" in response.headers["content-disposition"]


def test_экспорт_сводки_и_карточки_цфо(client, budget):
    assert client.get("/api/export/summary.xlsx").status_code == 200
    assert client.get(f"/api/export/cfo/{budget['cfo_id']}.xlsx").status_code == 200