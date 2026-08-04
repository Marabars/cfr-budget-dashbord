"""Общие утилиты чтения Excel: поиск шапки, сопоставление колонок, приведение типов."""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any, Iterable, Sequence

import openpyxl

from app.calc import MONTHS, MONTH_NAMES

MAX_HEADER_SCAN = 25   ***REMOVED*** в скольких первых строках ищем шапку
MAX_COLS = 60


def load_rows(source: str | Path | bytes, sheet: str | None = None,
              max_cols: int = MAX_COLS) -> tuple[list[list[Any]], str, list[str]]:
    """Читает лист в список строк. Возвращает (строки, имя листа, все листы)."""
    stream: Any = io.BytesIO(source) if isinstance(source, bytes) else source
    wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        name = sheet if sheet in names else names[0]
        ws = wb[name]
        rows = [list(r) for r in ws.iter_rows(max_col=max_cols, values_only=True)]
        return rows, name, names
    finally:
        wb.close()


def text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def num(value: Any) -> float:
    """Число из ячейки; пустое и нечисловое -> 0."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^\d,.\-]", "", str(value)).replace(",", ".")
    if cleaned in ("", "-", ".", "-."):
        return 0.0
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def iso_date(value: Any) -> str:
    """Дата в YYYY-MM-DD из даты Excel или строки; пустая строка при неудаче."""
    if value is None:
        return ""
    if hasattr(value, "year") and hasattr(value, "month"):
        return "%04d-%02d-%02d" % (value.year, value.month, getattr(value, "day", 1))
    s = str(value).strip()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return m.group(0)
    m = re.match(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})", s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y < 100:
            y += 2000
        return "%04d-%02d-%02d" % (y, mo, d)
    return ""


def _norm(s: str) -> str:
    return re.sub(r"[\s_\-–—.,:;()«»\"']+", " ", s.lower()).strip()


def find_header(rows: Sequence[Sequence[Any]], required: Iterable[str],
                limit: int = MAX_HEADER_SCAN) -> int:
    """Индекс строки шапки (0-based): первая строка, где встречены все обязательные слова."""
    needed = [_norm(r) for r in required]
    for i, row in enumerate(rows[:limit]):
        joined = " | ".join(_norm(text(c)) for c in row if c is not None)
        if all(n in joined for n in needed):
            return i
    return -1


def map_columns(header: Sequence[Any], aliases: dict[str, Sequence[str]]) -> dict[str, int]:
    """Сопоставляет поля с индексами колонок по списку допустимых названий.

    Два прохода: сначала точные совпадения, потом частичные. Иначе, например,
    алиас «бюджет» занял бы колонку «Бюджет - опер план» раньше колонки «БЮДЖЕТ».
    """
    cells = [(i, _norm(text(c))) for i, c in enumerate(header) if text(c)]
    mapping: dict[str, int] = {}
    used: set[int] = set()

    for exact in (True, False):
        for field, names in aliases.items():
            if field in mapping:
                continue
            for alias in names:
                a = _norm(alias)
                for idx, value in cells:
                    if idx in used or not value:
                        continue
                    if value == a if exact else a in value:
                        mapping[field] = idx
                        used.add(idx)
                        break
                if field in mapping:
                    break
    return mapping


def month_columns(header: Sequence[Any]) -> dict[int, int]:
    """Колонки месяцев: {1..12 -> индекс колонки}. Понимает названия и даты."""
    out: dict[int, int] = {}
    for idx, cell in enumerate(header):
        if cell is None:
            continue
        if hasattr(cell, "month") and hasattr(cell, "year"):
            out.setdefault(int(cell.month), idx)
            continue
        value = _norm(text(cell))
        for m, name in enumerate(MONTH_NAMES, start=1):
            if value.startswith(_norm(name)[:4]) and m not in out:
                out[m] = idx
                break
    return out if len(out) == MONTHS else out


def read_months(row: Sequence[Any], cols: dict[int, int]) -> list[float]:
    """Суммы по месяцам из строки согласно карте колонок."""
    out = [0.0] * MONTHS
    for m, idx in cols.items():
        if 1 <= m <= MONTHS and idx < len(row):
            out[m - 1] = num(row[idx])
    return out


def sample_values(rows: Sequence[Sequence[Any]], header_idx: int, col: int, count: int = 3) -> list[str]:
    """Примеры значений колонки — для экрана сопоставления колонок."""
    out: list[str] = []
    for row in rows[header_idx + 1:]:
        if col >= len(row):
            continue
        value = text(row[col])
        if value:
            out.append(value[:60])
        if len(out) >= count:
            break
    return out