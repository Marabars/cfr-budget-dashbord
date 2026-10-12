"""Экспорт в Excel. Суммы отдаются точными, с копейками."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from app import calc, service

MONEY = "***REMOVED*** ***REMOVED******REMOVED***0.00"

***REMOVED*** Палитра взята из файла ликвидности, а не придумана: выгрузку кладут рядом с
***REMOVED*** исходником и сверяют глазами, а лист «Оперплан ЦФО» прямо вставляют обратно
***REMOVED*** в книгу. Цвета прочитаны из самого файла — тема Office и tint уже разрешены
***REMOVED*** в hex, поэтому здесь плоские значения.
HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")       ***REMOVED*** шапка
TOTAL_FILL = PatternFill("solid", fgColor="D6DCE5")      ***REMOVED*** строка ИТОГО блока
ARTICLE_FILL = PatternFill("solid", fgColor="DCDCDC")    ***REMOVED*** № и «Статья оборотов»
TEXT_FILL = PatternFill("solid", fgColor="D6DCE5")       ***REMOVED*** ЦФО, контрагент, договор
NOTE_FILL = PatternFill("solid", fgColor="FFE4B5")       ***REMOVED*** «Комментарий»
OWN_NOTE_FILL = PatternFill("solid", fgColor="FBE5D6")   ***REMOVED*** «Комм наш»
MONTH_FILL = PatternFill("solid", fgColor="F0FFF0")      ***REMOVED*** месяцы
CALC_FILL = PatternFill("solid", fgColor="E5FFEB")       ***REMOVED*** расчётные колонки оперплана
MONEY_HEAD_FILL = PatternFill("solid", fgColor="DDF2FF")  ***REMOVED*** шапки денежных колонок
PLAN_HEAD_FILL = PatternFill("solid", fgColor="DDDDFF")   ***REMOVED*** шапка «БЮДЖЕТ»
REMAIN_HEAD_FILL = PatternFill("solid", fgColor="FFE7E7")  ***REMOVED*** шапка «Остаток бюджета»
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(bottom=THIN)


def _style_header(ws, row: int, last_col: int) -> None:
    for col in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True, size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _autosize(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _paint(ws, row: int, fills: dict[int, PatternFill]) -> None:
    """Заливка строки по карте «колонка → цвет». Колонки без цвета остаются белыми."""
    for col, fill in fills.items():
        ws.cell(row=row, column=col).fill = fill


def _autofilter(ws, header_row: int, last_col: int, last_row: int) -> None:
    """Кнопки фильтра на шапке — как в исходнике.

    Критериев не задаём, строки не прячутся: загрузка читает лист целиком и
    фильтр ей не мешает.
    """
    first = get_column_letter(1)
    last = get_column_letter(last_col)
    ws.auto_filter.ref = f"{first}{header_row}:{last}{max(last_row, header_row)}"


def _to_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


***REMOVED*** Раскладка листа «Бюджеты ЦФО» в одном месте: выгрузка и шаблон обязаны
***REMOVED*** совпадать колонка в колонку, иначе заполненный шаблон не зальётся.
***REMOVED*** Классификатор стоит сразу за ЦФО — он про ту же строку, что и дирекция.
BUDGET_HEADER = ["Статья оборотов", "ЦФО", "Классификатор", "Контрагент", "Договор",
                 "Детализация", "Комментарий", "ИТОГО"] + calc.MONTH_NAMES
COL = {title: i + 1 for i, title in enumerate(BUDGET_HEADER)}
MONTH_COL = COL["ИТОГО"] + 1           ***REMOVED*** первая колонка месяца
BUDGET_WIDTHS = {COL["Статья оборотов"]: 38, COL["ЦФО"]: 30, COL["Классификатор"]: 15,
                 COL["Контрагент"]: 28, COL["Договор"]: 26, COL["Детализация"]: 34,
                 COL["Комментарий"]: 30, COL["ИТОГО"]: 16}
***REMOVED*** Цвета колонок — оттуда же, откуда раскладка. «ИТОГО» в исходнике без заливки,
***REMOVED*** классификатор — наша колонка, встаёт в один ряд с ЦФО и контрагентом.
BUDGET_FILLS = {COL["Статья оборотов"]: ARTICLE_FILL, COL["ЦФО"]: TEXT_FILL,
                COL["Классификатор"]: TEXT_FILL, COL["Контрагент"]: TEXT_FILL,
                COL["Договор"]: TEXT_FILL, COL["Детализация"]: TEXT_FILL,
                COL["Комментарий"]: NOTE_FILL}
BUDGET_FILLS.update({MONTH_COL + m: MONTH_FILL for m in range(calc.MONTHS)})
LAST_BUDGET_COL = len(BUDGET_HEADER)


def budget_sheet(year: int, cfo_ids=None) -> bytes:
    """Лист «Бюджеты ЦФО» слоя планирования — тот же вид, что у исходного файла.

    Раскладка один в один повторяет то, что читает importers.cfo_budgets:
    блок ЦФО открывается строкой «ИТОГО» с контрольной суммой, под ней строки.
    Поэтому выгруженный файл можно поправить в Excel и залить обратно.
    """
    return _budget_layout_sheet(year, cfo_ids, fact=False)


def fact_sheet(year: int, cfo_ids=None) -> bytes:
    """Факт года той же раскладкой: в ИТОГО и месяцах стоит факт, а не план.

    Лист называется так же и читается той же загрузкой, что и бюджет, —
    круг «выгрузили факт → поправили → залили» замкнут. Строки плана, по
    которым факта нет, остаются с нулями: загрузка факта их пропускает.
    """
    return _budget_layout_sheet(year, cfo_ids, fact=True)


def _budget_layout_sheet(year: int, cfo_ids=None, fact: bool = False) -> bytes:
    """Общая раскладка листа «Бюджеты ЦФО»: план года или факт года."""
    from app import operplan as operplan_service

    view = operplan_service.view(year, cfo_ids)
    facts = (operplan_service.fact_months_by_row(int(view["year"]), cfo_ids)
             if fact else None)
    wb = Workbook()
    ws = wb.active
    ws.title = cfo_budgets_sheet_name()

    header = BUDGET_HEADER
    title = (f"Факт {view['year']} года по строкам бюджета" if fact
             else f"Бюджеты ЦФО, {view['year']} год")
    ws.cell(row=1, column=1, value=title)
    ws.cell(row=1, column=1).font = Font(bold=True, size=12)
    for col, title_cell in enumerate(header, start=1):
        ws.cell(row=3, column=col, value=title_cell)
    _style_header(ws, 3, len(header))

    def money(r: int, c: int, value: float, bold: bool = False) -> None:
        cell = ws.cell(row=r, column=c, value=round(float(value or 0), 2))
        cell.number_format = MONEY
        if bold:
            cell.font = Font(bold=True)

    def amounts(row: dict) -> tuple[float, list[float]]:
        """Суммы строки: план из вида или факт по месяцам, ИТОГО — их сумма."""
        if facts is None:
            return row["plan_year"], row["months"]
        months = facts.get(int(row["id"])) or [0.0] * calc.MONTHS
        return sum(months), months

    idx = 4
    for block in view["blocks"]:
        rows = [(row, *amounts(row)) for row in block["rows"]]
        block_total = (sum(total for _, total, _ in rows) if facts is not None
                       else block["plan_year"])
        block_months = ([sum(m[i] for _, _, m in rows) for i in range(calc.MONTHS)]
                        if facts is not None else block["months"])

        ws.cell(row=idx, column=1, value="ИТОГО").font = Font(bold=True)
        ws.cell(row=idx, column=2, value=block["cfo_name"]).font = Font(bold=True)
        ws.cell(row=idx, column=COL["Договор"], value=block["budget_name"])
        money(idx, COL["ИТОГО"], block_total, bold=True)
        for m, amount in enumerate(block_months):
            money(idx, MONTH_COL + m, amount, bold=True)
        for col in range(1, len(header) + 1):
            ws.cell(row=idx, column=col).fill = TOTAL_FILL
        idx += 1

        for row, total, months in rows:
            _paint(ws, idx, BUDGET_FILLS)
            ws.cell(row=idx, column=COL["Статья оборотов"], value=row["article"])
            ws.cell(row=idx, column=COL["ЦФО"], value=block["cfo_name"])
            ws.cell(row=idx, column=COL["Классификатор"], value=row["classifier"])
            ws.cell(row=idx, column=COL["Контрагент"], value=row["counterparty"])
            ws.cell(row=idx, column=COL["Договор"], value=row["contract"])
            ws.cell(row=idx, column=COL["Детализация"], value=row["detail"])
            ws.cell(row=idx, column=COL["Комментарий"], value=row["comment"])
            money(idx, COL["ИТОГО"], total)
            for m, amount in enumerate(months):
                money(idx, MONTH_COL + m, amount)
            idx += 1

    _autofilter(ws, 3, LAST_BUDGET_COL, idx - 1)
    _autosize(ws, BUDGET_WIDTHS)
    for m in range(calc.MONTHS):
        ws.column_dimensions[get_column_letter(MONTH_COL + m)].width = 14
    return _to_bytes(wb)


def cfo_budgets_sheet_name() -> str:
    from app.importers import cfo_budgets
    return cfo_budgets.SHEET_NAME


TEMPLATE_ROWS = 12          ***REMOVED*** пустых строк на блок — хватает на один ЦФО без вставок
REFS_SHEET = "Справочник"
REFS_COL = {"Статьи оборотов": 2, "Контрагенты": 3}
ALL_CFO = 0                 ***REMOVED*** ключ общего блока справочника: у ЦФО id идут с единицы


def budget_template(year: int, cfo_ids=None) -> bytes:
    """Пустой лист «Бюджеты ЦФО» под заполнение руками.

    Раскладка та же, что у `budget_sheet`, поэтому заполненный шаблон принимает
    та же загрузка. Блок каждого ЦФО открыт строкой «ИТОГО» — держателем бюджета
    считается блок, а не колонка «ЦФО» в самой строке, так что строки достаточно
    вписывать под нужным заголовком.

    Пустые строки при загрузке пропускаются: лишние заготовки можно не удалять.
    """
    from app import operplan as operplan_service, repo
    from app.importers.cfo_budgets import DEFAULT_BUDGET

    allowed = None if cfo_ids is None else {int(i) for i in cfo_ids}
    cfo_list = [c for c in repo.list_cfo()
                if allowed is None or int(c["id"]) in allowed]
    budgets_by_cfo: dict[int, list[str]] = {}
    for block in operplan_service.view(year, cfo_ids)["blocks"]:
        budgets_by_cfo.setdefault(int(block["cfo_id"]), []).append(block["budget_name"])

    wb = Workbook()
    ws = wb.active
    ws.title = cfo_budgets_sheet_name()

    header = BUDGET_HEADER
    ws.cell(row=1, column=1, value=f"Бюджеты ЦФО, {year} год — шаблон для заполнения")
    ws.cell(row=1, column=1).font = Font(bold=True, size=12)
    for col, title in enumerate(header, start=1):
        ws.cell(row=3, column=col, value=title)
    _style_header(ws, 3, len(header))
    _total_note(ws, 3)

    idx = 4
    blocks: list[tuple[int, int, int]] = []     ***REMOVED*** ЦФО и диапазон его пустых строк
    for cfo in cfo_list:
        names = budgets_by_cfo.get(int(cfo["id"])) or [DEFAULT_BUDGET]
        for budget_name in names:
            ws.cell(row=idx, column=1, value="ИТОГО").font = Font(bold=True)
            ws.cell(row=idx, column=2, value=cfo["name"]).font = Font(bold=True)
            ***REMOVED*** В колонке «Договор» строки ИТОГО лежит название бюджета блока,
            ***REMOVED*** а не договор: загрузка читает его оттуда. Пустая клетка значит
            ***REMOVED*** «Основной бюджет», поэтому у дирекции с единственным бюджетом её
            ***REMOVED*** не заполняем — в пустом шаблоне это читалось как настоящий договор.
            if budget_name != DEFAULT_BUDGET:
                ws.cell(row=idx, column=COL["Договор"], value=budget_name)
            for col in range(1, len(header) + 1):
                ws.cell(row=idx, column=col).fill = TOTAL_FILL
            idx += 1
            first = idx
            for _ in range(TEMPLATE_ROWS):
                _paint(ws, idx, BUDGET_FILLS)
                cell = ws.cell(row=idx, column=COL["ИТОГО"])
                cell.number_format = MONEY
                for m in range(calc.MONTHS):
                    ws.cell(row=idx, column=MONTH_COL + m).number_format = MONEY
                idx += 1
            blocks.append((int(cfo["id"]), first, idx - 1))

    _classifier_choice(ws, 4, idx - 1)
    _autofilter(ws, 3, LAST_BUDGET_COL, idx - 1)
    _autosize(ws, BUDGET_WIDTHS)
    for m in range(calc.MONTHS):
        ws.column_dimensions[get_column_letter(MONTH_COL + m)].width = 14

    ranges = _template_refs(wb, year, cfo_ids,
                            [cfo_id for cfo_id, _, _ in blocks])
    _reference_choice(wb, ws, blocks, ranges)
    return _to_bytes(wb)


def _classifier_choice(ws, first_row: int, last_row: int) -> None:
    """Выпадающий список МРГ/МРК в колонке классификатора.

    Значения ровно два, поэтому проще дать выбрать, чем ловить опечатку потом.
    Список — подсказка, а не запрет: загрузка всё равно приводит чужое значение
    к МРГ, иначе строка бюджета терялась бы из-за описки в одной колонке.
    """
    if last_row < first_row:
        return
    letter = get_column_letter(COL["Классификатор"])
    rule = DataValidation(type="list", formula1='"МРГ,МРК"', allow_blank=True,
                          showDropDown=False)
    rule.error = "Допустимы только МРГ и МРК"
    rule.prompt = "Выберите МРГ или МРК"
    ws.add_data_validation(rule)
    rule.add(f"{letter}{first_row}:{letter}{last_row}")


def _total_note(ws, row: int) -> None:
    """Примечание на шапке колонки ИТОГО.

    Единственная подсказка, оставшаяся в шаблоне: отдельный лист «Как заполнять»
    убран, а это правило не выводится из самого листа — пустое ИТОГО при
    заполненных месяцах выглядит как потерянная сумма.
    """
    cell = ws.cell(row=row, column=COL["ИТОГО"])
    cell.comment = Comment(
        "Если заполнены только месяцы, ИТОГО будет посчитано как их сумма.",
        "Бюджет ЦФО", height=70, width=260)


def _template_refs(wb: Workbook, year: int, cfo_ids=None,
                   block_cfo_ids: list[int] | None = None) -> dict[int, dict[str, str]]:
    """Лист-справочник: какие статьи и контрагенты уже встречались у ЦФО.

    Отдаёт имена диапазонов для выпадающих списков: `{cfo_id: {колонка: имя}}`,
    плюс — когда он кому-то нужен — общий блок под ключом `ALL_CFO`.

    Нужен, чтобы в шаблоне писали существующие названия, а не новые варианты
    написания того же: статья с опечаткой заведёт в плане отдельную строку.

    Общий блок — запасной список для дирекции, у которой своего ещё нет. Если
    своё есть у всех блоков шаблона, он на лист не попадает: при выгрузке одной
    дирекции он слово в слово повторял её собственный список.
    """
    from app import operplan as operplan_service, repo

    refs = operplan_service.references(year, None if cfo_ids is None else list(cfo_ids))
    names = {int(c["id"]): c["name"] for c in repo.list_cfo()}
    own = {int(k): v for k, v in refs["by_cfo"].items()}
    needs_all = any(not (own.get(int(cfo_id)) or {}).get(column)
                    for cfo_id in (block_cfo_ids or [])
                    for column in ("articles", "counterparties"))
    ws = wb.create_sheet(REFS_SHEET)
    for col, title in enumerate(["ЦФО", "Статьи оборотов", "Контрагенты"], start=1):
        ws.cell(row=1, column=col, value=title)
    _style_header(ws, 1, 3)

    out: dict[int, dict[str, str]] = {}

    def group(label: str, key: str, articles: list[str], counterparties: list[str],
              idx: int) -> int:
        """Один блок справочника; возвращает строку, с которой начинается следующий."""
        height = max(len(articles), len(counterparties))
        if not height:
            return idx
        ws.cell(row=idx, column=1, value=label).font = Font(bold=True)
        for offset in range(height):
            if offset < len(articles):
                ws.cell(row=idx + offset, column=2, value=articles[offset])
            if offset < len(counterparties):
                ws.cell(row=idx + offset, column=3, value=counterparties[offset])
        bucket: dict[str, str] = {}
        for column, values, prefix in (("Статьи оборотов", articles, "_art"),
                                       ("Контрагенты", counterparties, "_cp")):
            if not values:
                continue
            bucket[column] = _define_range(wb, f"{prefix}{key}", ws.title, column,
                                           idx, idx + len(values) - 1)
        if bucket:
            out[key] = bucket
        return idx + height + 1

    idx = 2
    if needs_all:
        idx = group("— все дирекции —", ALL_CFO, refs["all"]["articles"],
                    refs["all"]["counterparties"], idx)
    for cfo_id, bucket in sorted(refs["by_cfo"].items(),
                                 key=lambda kv: names.get(int(kv[0]), "")):
        idx = group(names.get(int(cfo_id), ""), int(cfo_id),
                    bucket["articles"], bucket["counterparties"], idx)
    _autosize(ws, {1: 30, 2: 44, 3: 34})
    return out


def _define_range(wb: Workbook, name: str, sheet: str, column: str,
                  first_row: int, last_row: int) -> str:
    """Именованный диапазон на листе справочника.

    Проверка данных ссылается на имя, а не на `'Справочник'!$B$2:$B$9` напрямую:
    ссылку на другой лист внутри проверки понимают не все версии Excel, имя —
    любая.
    """
    letter = get_column_letter(REFS_COL[column])
    ref = f"'{sheet}'!${letter}${first_row}:${letter}${last_row}"
    wb.defined_names[name] = DefinedName(name, attr_text=ref)
    return name


def _reference_choice(wb: Workbook, ws, blocks: list[tuple[int, int, int]],
                      ranges: dict[int, dict[str, str]]) -> None:
    """Выпадающие списки статей и контрагентов — свои у каждого блока ЦФО.

    Список берётся с листа «Справочник»: у дирекции свой набор, а если у неё
    там пусто (бюджета ещё не было) — общий, иначе заполнять пришлось бы вовсе
    без подсказки.

    Запрета нет (`showErrorMessage=False`): новая статья в плане законна, список
    лишь помогает не выдумать второе написание уже заведённой.
    """
    for column, target in (("Статья оборотов", "Статьи оборотов"),
                           ("Контрагент", "Контрагенты")):
        letter = get_column_letter(COL[column])
        for cfo_id, first, last in blocks:
            name = ((ranges.get(int(cfo_id)) or {}).get(target)
                    or (ranges.get(ALL_CFO) or {}).get(target))
            if not name or last < first:
                continue
            rule = DataValidation(type="list", formula1=name, allow_blank=True,
                                  showDropDown=False, showErrorMessage=False)
            ws.add_data_validation(rule)
            rule.add(f"{letter}{first}:{letter}{last}")


***REMOVED*** Раскладка «Оперплан ЦФО» своя: № впереди, за месяцами идут расчётные колонки.
***REMOVED*** Цвета — из того же листа исходника, туда выгрузку и вставляют.
OPERPLAN_FILLS = {1: ARTICLE_FILL, 2: ARTICLE_FILL, 3: TEXT_FILL, 4: TEXT_FILL,
                  5: TEXT_FILL, 6: TEXT_FILL, 7: NOTE_FILL, 8: OWN_NOTE_FILL,
                  10: HEAD_FILL, 11: CALC_FILL, 12: CALC_FILL, 13: CALC_FILL}
OPERPLAN_HEAD_FILLS = {8: OWN_NOTE_FILL, 9: MONEY_HEAD_FILL, 10: PLAN_HEAD_FILL,
                       11: CALC_FILL, 12: CALC_FILL, 13: CALC_FILL,
                       14: MONEY_HEAD_FILL, 15: REMAIN_HEAD_FILL}


def operplan(year: int | None = None, cfo_ids=None) -> bytes:
    """Лист «Оперплан ЦФО» в раскладке исходного файла — для вставки в «Ликвидность».

    Порядок колонок и служебные строки повторяют источник: строка с метками
    ФАКТ / ПРОГНОЗ, шапка, затем блоки «ИТОГО» с их строками.
    """
    view = service.year_view(year, cfo_ids)
    cfg = view["settings"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Оперплан ЦФО"

    first_month_col = 16
    for m in range(calc.MONTHS):
        ws.cell(row=1, column=first_month_col + m,
                value=f"{cfg['year']}-{m + 1:02d}-01")
        ws.cell(row=3, column=first_month_col + m,
                value="ФАКТ" if m + 1 <= cfg["closed_through"] else "ПРОГНОЗ").font = Font(bold=True)

    header = ["№", "Статья оборотов", "ЦФО", "Контрагент", "Договор", "Детализация",
              "Комментарий", "Комм наш", "Бюджет - опер план", "БЮДЖЕТ", "Оперплан фикс",
              "", "Итого опер план", "Факт", "Остаток бюджета"] + calc.MONTH_NAMES
    first_month_fills = {first_month_col + m: MONTH_FILL for m in range(calc.MONTHS)}
    body_fills = {**OPERPLAN_FILLS, **first_month_fills}
    for col, title in enumerate(header, start=1):
        ws.cell(row=5, column=col, value=title)
    _style_header(ws, 5, len(header))
    _paint(ws, 5, OPERPLAN_HEAD_FILLS)       ***REMOVED*** денежные колонки в шапке своего цвета

    row_idx = 6
    for block in view["blocks"]:
        budget = view["budgets"][block["id"]]
        t = block["totals"]
        ws.cell(row=row_idx, column=2, value="ИТОГО").font = Font(bold=True)
        ws.cell(row=row_idx, column=3, value=block["cfo_name"]).font = Font(bold=True)
        ws.cell(row=row_idx, column=5, value=block["name"])
        for col, value in ((9, t["free_limit"]), (10, t["plan_year"]),
                           (13, t["operplan"]), (14, t["fact"]), (15, t["cash_remain"])):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.number_format = MONEY
            cell.font = Font(bold=True)
        for m in range(calc.MONTHS):
            total_m = sum(r["fact_months"][m] + r["claim_months"][m] for r in budget["rows"])
            cell = ws.cell(row=row_idx, column=first_month_col + m, value=total_m)
            cell.number_format = MONEY
            cell.font = Font(bold=True)
        for col in range(1, len(header) + 1):
            ws.cell(row=row_idx, column=col).fill = TOTAL_FILL
        row_idx += 1

        for r in budget["rows"]:
            _paint(ws, row_idx, body_fills)
            values = [
                r["row_no"], r["article"], block["cfo_name"], r["counterparty"], r["contract"],
                r["detail"], r["comment"], r["approver_comment"],
                r["free_limit"], r["plan_year"], r["plan_year"], None,
                r["operplan"], r["fact"], r["cash_remain"],
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                if col >= 9 and isinstance(value, (int, float)):
                    cell.number_format = MONEY
            for m in range(calc.MONTHS):
                cell = ws.cell(row=row_idx, column=first_month_col + m,
                               value=r["fact_months"][m] + r["claim_months"][m])
                cell.number_format = MONEY
            row_idx += 1

    _autofilter(ws, 5, len(header), row_idx - 1)
    _autosize(ws, {1: 6, 2: 38, 3: 34, 4: 26, 5: 26, 6: 26, 7: 26, 8: 20,
                   9: 18, 10: 18, 11: 16, 12: 3, 13: 18, 14: 16, 15: 18})
    for m in range(calc.MONTHS):
        ws.column_dimensions[get_column_letter(first_month_col + m)].width = 15
    return _to_bytes(wb)


def summary(year: int | None = None, cfo_ids=None) -> bytes:
    """Сводка по ЦФО: план / факт / заявки / остаток / статус."""
    view = service.year_view(year, cfo_ids)
    cfg = view["settings"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка по ЦФО"

    ws.cell(row=1, column=1, value=f"Бюджеты ЦФО {cfg['year']} · на {cfg['as_of']} · "
                                   f"факт закрыт по {calc.MONTH_NAMES[cfg['closed_through'] - 1] if cfg['closed_through'] else '—'}")
    ws.cell(row=1, column=1).font = Font(bold=True, size=12)

    header = ["ЦФО", "Бюджет", "Бюджет год", "Факт", "Заявки", "Итого опер план",
              "Остаток бюджета", "Свободный лимит", "Освоение, %", "Статус", "Вне бюджета"]
    for col, title in enumerate(header, start=1):
        ws.cell(row=3, column=col, value=title)
    _style_header(ws, 3, len(header))

    row_idx = 4
    for b in view["blocks"]:
        t = b["totals"]
        values = [b["cfo_name"], b["name"], t["plan_year"], t["fact"], t["claims"],
                  t["operplan"], t["cash_remain"], t["free_limit"], t["pct"],
                  b["status_text"], b["out_of_budget"]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            if col in (3, 4, 5, 6, 7, 8, 11):
                cell.number_format = MONEY
            if col == 9:
                cell.number_format = "0.0%"
        row_idx += 1

    totals_row = row_idx
    ws.cell(row=totals_row, column=1, value="ИТОГО").font = Font(bold=True)
    for col in (3, 4, 5, 6, 7, 8, 11):
        letter = get_column_letter(col)
        cell = ws.cell(row=totals_row, column=col,
                       value=f"=SUM({letter}4:{letter}{row_idx - 1})")
        cell.number_format = MONEY
        cell.font = Font(bold=True)

    ***REMOVED*** строка ИТОГО в фильтр не входит: сортировкой её утащило бы в середину
    _autofilter(ws, 3, len(header), totals_row - 1)
    _autosize(ws, {1: 46, 2: 28, 3: 18, 4: 18, 5: 16, 6: 18, 7: 18, 8: 18, 9: 12, 10: 22, 11: 16})
    return _to_bytes(wb)


def cfo_card(cfo_id: int, year: int | None = None) -> bytes:
    """Карточка одного ЦФО: строки с месяцами."""
    card = service.cfo_card(cfo_id, year)
    wb = Workbook()
    ws = wb.active
    ws.title = (card["cfo"]["name"] if card["cfo"] else "ЦФО")[:31]

    header = ["№", "Статья оборотов", "Контрагент", "Договор", "Детализация",
              "Бюджет", "Итого опер план", "Факт", "Остаток бюджета",
              "Свободный лимит"] + calc.MONTH_NAMES
    card_fills = {1: ARTICLE_FILL, 2: ARTICLE_FILL, 3: TEXT_FILL, 4: TEXT_FILL,
                  5: TEXT_FILL}
    card_fills.update({11 + m: MONTH_FILL for m in range(calc.MONTHS)})
    for col, title in enumerate(header, start=1):
        ws.cell(row=1, column=col, value=title)
    _style_header(ws, 1, len(header))

    row_idx = 2
    for budget in card["budgets"]:
        ws.cell(row=row_idx, column=1, value=budget["name"]).font = Font(bold=True)
        for col in range(1, len(header) + 1):
            ws.cell(row=row_idx, column=col).fill = TOTAL_FILL
        row_idx += 1
        for r in budget["rows"]:
            _paint(ws, row_idx, card_fills)
            values = [r["row_no"], r["article"], r["counterparty"], r["contract"], r["detail"],
                      r["plan_year"], r["operplan"], r["fact"], r["cash_remain"], r["free_limit"]]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                if col >= 6:
                    cell.number_format = MONEY
            for m in range(calc.MONTHS):
                cell = ws.cell(row=row_idx, column=11 + m,
                               value=r["fact_months"][m] + r["claim_months"][m])
                cell.number_format = MONEY
            row_idx += 1

    _autofilter(ws, 1, len(header), row_idx - 1)
    _autosize(ws, {1: 6, 2: 38, 3: 26, 4: 26, 5: 30, 6: 16, 7: 18, 8: 16, 9: 18, 10: 18})
    return _to_bytes(wb)


def out_of_budget(year: int | None = None) -> bytes:
    """Реестр расходов вне бюджета."""
    from app import repo

    cfg = service.settings()
    items = repo.list_out_of_budget(year or cfg["year"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Вне бюджета"
    header = ["Дата", "ЦФО", "Контрагент", "Договор", "Назначение", "Сумма", "Кто пометил"]
    for col, title in enumerate(header, start=1):
        ws.cell(row=1, column=col, value=title)
    _style_header(ws, 1, len(header))
    for i, f in enumerate(items, start=2):
        values = [f["pay_date"], f["cfo_name"] or "", f["counterparty"], f["contract"],
                  f["purpose"], float(f["amount"]), f["matched_by"]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            if col == 6:
                cell.number_format = MONEY
    _autofilter(ws, 1, len(header), len(items) + 1)
    _autosize(ws, {1: 12, 2: 40, 3: 30, 4: 26, 5: 40, 6: 16, 7: 24})
    return _to_bytes(wb)