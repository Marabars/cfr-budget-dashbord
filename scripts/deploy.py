# -*- coding: utf-8 -*-
"""Деплой приложения на сервер по SSH.

Секретов в коде нет: доступ берётся из переменных окружения или из внешнего
xlsx-файла, где в первой колонке имя параметра (login / pass / host или vm),
а во второй — значение. Пароль не печатается и не пишется в логи.

    set DEPLOY_CREDS=<путь к файлу.xlsx>     # или DEPLOY_USER / DEPLOY_PASS / DEPLOY_HOST
    python scripts/deploy.py check           # что на сервере: python, порт, служба
    python scripts/deploy.py deploy          # выложить код и перезапустить
    python scripts/deploy.py restart         # только перезапустить
    python scripts/deploy.py logs            # последние строки журнала
    python scripts/deploy.py stop            # остановить

Автозапуск: если для пользователя удаётся включить linger, приложение
регистрируется как служба пользовательского systemd (поднимается после
перезагрузки сервера). Иначе запускается через nohup — как соседние
приложения на этом сервере.
"""
from __future__ import annotations

import os
import sys
import tarfile
import tempfile
from pathlib import Path

try:
    import paramiko
except ImportError:  # pragma: no cover
    print("Нужен paramiko:  pip install paramiko")
    raise SystemExit(2)

ROOT = Path(__file__).resolve().parent.parent
APP_DIR = "cfo-budget-dashboard"          # каталог в $HOME
SERVICE = "cfo-budget"
PORT = int(os.environ.get("DEPLOY_PORT", "8000"))

INCLUDE = ["app", "static", "scripts", "deploy", "requirements.txt", "README.md"]
# instruction — генератор инструкции для дирекций: демо-база и 2,4 МБ снимков,
# приложению они не нужны и на сервере только занимают место
EXCLUDE_DIRS = {"__pycache__", ".git", ".pytest_cache", "data", "venv", ".venv",
                "instruction"}
EXCLUDE_SUFFIX = {".pyc", ".db", ".xlsx", ".xls", ".csv", ".log"}


# ---------- доступ ----------

def credentials() -> tuple[str, str, str]:
    """Логин, пароль и хост из окружения или из xlsx с данными сервера."""
    user = os.environ.get("DEPLOY_USER", "")
    password = os.environ.get("DEPLOY_PASS", "")
    host = os.environ.get("DEPLOY_HOST", "")

    creds_file = os.environ.get("DEPLOY_CREDS", "")
    if creds_file and (not user or not password or not host):
        import openpyxl

        wb = openpyxl.load_workbook(creds_file, read_only=True, data_only=True)
        try:
            for sheet in wb.sheetnames:
                for row in wb[sheet].iter_rows(max_row=50, max_col=6, values_only=True):
                    cells = [("" if c is None else str(c).strip()) for c in row]
                    if len(cells) < 2 or not cells[0] or not cells[1]:
                        continue
                    key, value = cells[0].lower(), cells[1]
                    if "login" in key and not user:
                        user = value
                    elif "pass" in key and not password:
                        password = value
                    elif ("vm" in key or "host" in key) and not host:
                        host = value
        finally:
            wb.close()

    if not (user and password and host):
        print("Не хватает данных доступа: задайте DEPLOY_CREDS (путь к xlsx) "
              "или DEPLOY_USER / DEPLOY_PASS / DEPLOY_HOST.")
        raise SystemExit(2)
    return user, password, host


def connect() -> tuple[paramiko.SSHClient, str, str]:
    """Подключение по SSH. Пароль возвращается только для sudo и не логируется."""
    user, password, host = credentials()
    client = paramiko.SSHClient()
    client.load_system_host_keys()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    client.connect(hostname=host, username=user, password=password, timeout=30,
                   look_for_keys=False, allow_agent=False)
    return client, host, password


def run(client: paramiko.SSHClient, command: str, check: bool = True,
        quiet: bool = False) -> tuple[int, str]:
    _, stdout, stderr = client.exec_command(command, timeout=900)
    out = stdout.read().decode("utf-8", "replace")
    err = stderr.read().decode("utf-8", "replace")
    code = stdout.channel.recv_exit_status()
    text = (out + err).strip()
    if not quiet and text:
        print(text)
    if check and code != 0:
        raise SystemExit(f"Команда завершилась с кодом {code}: {command}")
    return code, text


def sudo(client: paramiko.SSHClient, password: str, command: str,
         check: bool = False, quiet: bool = True) -> tuple[int, str]:
    """Команда через sudo; пароль идёт в stdin, поэтому не виден в `ps`."""
    stdin, stdout, stderr = client.exec_command(f"sudo -S -p '' {command}", timeout=900)
    stdin.write(password + "\n")
    stdin.flush()
    out = stdout.read().decode("utf-8", "replace")
    err = stderr.read().decode("utf-8", "replace")
    code = stdout.channel.recv_exit_status()
    text = (out + err).replace("[sudo] password for", "").strip()
    if not quiet and text:
        print(text)
    if check and code != 0:
        raise SystemExit(f"sudo завершился с кодом {code}: {command}")
    return code, text


# ---------- упаковка ----------

def make_archive() -> Path:
    """Складывает код в tar.gz, исключая данные, венвы и служебное."""
    tmp = Path(tempfile.gettempdir()) / f"{APP_DIR}.tar.gz"

    def keep(rel: Path) -> bool:
        if any(part in EXCLUDE_DIRS for part in rel.parts):
            return False
        return rel.suffix.lower() not in EXCLUDE_SUFFIX

    with tarfile.open(tmp, "w:gz") as tar:
        for item in INCLUDE:
            src = ROOT / item
            if not src.exists():
                continue
            if src.is_file():
                if keep(Path(item)):
                    tar.add(src, arcname=item)
                continue
            for file in sorted(src.rglob("*")):
                rel = file.relative_to(ROOT)
                if file.is_file() and keep(rel):
                    tar.add(file, arcname=str(rel).replace("\\", "/"))
    return tmp


# ---------- запуск на сервере ----------

def user_unit(target: str) -> str:
    return f"""[Unit]
Description=Дашборд бюджетов ЦФО
After=network.target

[Service]
Type=simple
WorkingDirectory={target}
Environment=BUDGET_DB={target}/data/budget.db
ExecStart={target}/venv/bin/python -m uvicorn app.main:app --host 0.0.0.0 --port {PORT}
Restart=always
RestartSec=5

[Install]
WantedBy=default.target
"""


def linger_enabled(client: paramiko.SSHClient) -> bool:
    return "Linger=yes" in run(client, "loginctl show-user $(whoami)",
                               check=False, quiet=True)[1]


def enable_linger(client: paramiko.SSHClient, password: str) -> bool:
    """Linger нужен, чтобы пользовательские службы жили без активной сессии."""
    if linger_enabled(client):
        return True
    run(client, "loginctl enable-linger $(whoami)", check=False, quiet=True)
    if linger_enabled(client):
        return True
    sudo(client, password, "loginctl enable-linger $(whoami)")
    return linger_enabled(client)


def start_service(client: paramiko.SSHClient, password: str, target: str, user: str) -> str:
    """Ставит службу пользовательского systemd, при неудаче — nohup."""
    if enable_linger(client, password):
        run(client, "mkdir -p ~/.config/systemd/user")
        sftp = client.open_sftp()
        try:
            with sftp.open(f"/home/{user}/.config/systemd/user/{SERVICE}.service", "w") as fh:
                fh.write(user_unit(target))
        finally:
            sftp.close()
        code, _ = run(client, f"systemctl --user daemon-reload && "
                              f"systemctl --user enable {SERVICE} && "
                              f"systemctl --user restart {SERVICE}", check=False, quiet=True)
        if code == 0:
            return "systemd пользователя — поднимется после перезагрузки сервера"

    run(client, f"pkill -f 'uvicorn app.main:app.*{PORT}' || true", check=False, quiet=True)
    run(client, f"cd {target} && BUDGET_DB={target}/data/budget.db nohup "
                f"./venv/bin/python -m uvicorn app.main:app --host 0.0.0.0 --port {PORT} "
                f"> {target}/data/app.log 2>&1 & sleep 1", check=False, quiet=True)
    return "nohup — после перезагрузки сервера нужно запускать заново"


def service_active(client: paramiko.SSHClient) -> bool:
    return run(client, f"systemctl --user is-active {SERVICE}",
               check=False, quiet=True)[1].strip() == "active"


# ---------- команды ----------

def cmd_check() -> None:
    client, host, password = connect()
    try:
        print(f"=== {host} ===")
        run(client, "hostname; whoami; uname -sr")
        print("\n--- python ---")
        run(client, "python3 --version", check=False)
        print("\n--- sudo ---")
        if run(client, "sudo -n true", check=False, quiet=True)[0] == 0:
            print("  доступен без пароля")
        elif sudo(client, password, "-v true")[0] == 0:
            print("  доступен по паролю")
        else:
            print("  недоступен — используем пользовательский systemd или nohup")
        print("\n--- linger (автозапуск пользовательских служб) ---")
        run(client, "loginctl show-user $(whoami) | grep -i linger", check=False)
        print(f"\n--- порт {PORT} ---")
        _, text = run(client, f"(ss -ltn 2>/dev/null || netstat -ltn) | grep ':{PORT} ' || true",
                      check=False, quiet=True)
        print(f"  занят: {text}" if text else "  свободен")
        print(f"\n--- служба {SERVICE} ---")
        run(client, f"systemctl --user is-active {SERVICE} 2>/dev/null || echo '  не установлена'",
            check=False)
        print("\n--- приложение ---")
        run(client, f"curl -fsS http://127.0.0.1:{PORT}/health || echo '  не отвечает'", check=False)
    finally:
        client.close()


def cmd_deploy() -> None:
    user, _, _ = credentials()
    client, host, password = connect()
    try:
        target = f"/home/{user}/{APP_DIR}"
        print(f"каталог: {target}   порт: {PORT}")

        archive = make_archive()
        print(f"архив: {archive.name} ({archive.stat().st_size // 1024} КБ)")

        run(client, f"mkdir -p {target}/data")
        sftp = client.open_sftp()
        remote = f"/tmp/{APP_DIR}.tar.gz"
        try:
            sftp.put(str(archive), remote)
        finally:
            sftp.close()
        print("файлы загружены")

        # код обновляется, база в data/ остаётся на месте
        run(client, f"cd {target} && tar xzf {remote} && rm -f {remote}")

        print("\n--- venv и зависимости ---")
        run(client, f"cd {target} && (test -d venv || python3 -m venv venv) && "
                    f"./venv/bin/pip install -q --upgrade pip && "
                    f"./venv/bin/pip install -q -r requirements.txt && echo '  готово'")

        print("\n--- запуск ---")
        print("  " + start_service(client, password, target, user))

        print("\n--- проверка ---")
        run(client, f"sleep 5; curl -fsS http://127.0.0.1:{PORT}/health || "
                    f"echo '  приложение не ответило, смотрите logs'", check=False)
        print(f"\nГотово: http://{host}:{PORT}")
    finally:
        client.close()


def cmd_restart() -> None:
    user, _, _ = credentials()
    client, _, password = connect()
    try:
        target = f"/home/{user}/{APP_DIR}"
        print(start_service(client, password, target, user))
        run(client, f"sleep 4; curl -fsS http://127.0.0.1:{PORT}/health", check=False)
    finally:
        client.close()


def cmd_logs() -> None:
    user, _, _ = credentials()
    client, _, _ = connect()
    try:
        if service_active(client):
            run(client, f"journalctl --user -u {SERVICE} -n 40 --no-pager", check=False)
        else:
            run(client, f"tail -n 40 /home/{user}/{APP_DIR}/data/app.log", check=False)
    finally:
        client.close()


def cmd_stop() -> None:
    client, _, _ = connect()
    try:
        run(client, f"systemctl --user stop {SERVICE} 2>/dev/null; "
                    f"pkill -f 'uvicorn app.main:app.*{PORT}' || true", check=False)
        print("остановлено")
    finally:
        client.close()


COMMANDS = {"check": cmd_check, "deploy": cmd_deploy, "restart": cmd_restart,
            "logs": cmd_logs, "stop": cmd_stop}


def main() -> int:
    command = sys.argv[1] if len(sys.argv) > 1 else "check"
    handler = COMMANDS.get(command)
    if not handler:
        print(f"Команды: {', '.join(COMMANDS)}")
        return 2
    handler()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())